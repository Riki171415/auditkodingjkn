import os
import json
import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE

def add_heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT if level > 1 else WD_ALIGN_PARAGRAPH.CENTER
    for run in h.runs:
        run.font.name = 'Times New Roman'
        run.font.color.rgb = RGBColor(0, 0, 0)
    return h

def main():
    from modules.report_data import load_snapshot
    from modules.report_word import write_report
    snapshot = load_snapshot()
    return write_report(snapshot['cases'], 'exports/Laporan_Akhir_Nasional_Agregat.docx', snapshot['snapshot_id'], snapshot['hospitals'])


def _main_legacy():
    # 1. Read Markdown content
    md_path = 'exports/Laporan_Akhir_Nasional_Agregat.md'
    if not os.path.exists(md_path):
        print("MD file not found!")
        return
        
    with open(md_path, 'r', encoding='utf-8') as f:
        md_text = f.read()

    doc = Document()
    
    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    
    # Parse and write Markdown
    lines = md_text.split('\n')
    for line in lines:
        line = line.strip()
        if not line:
            continue
            
        if line.startswith('# '):
            add_heading(doc, line[2:].strip(), 1)
        elif line.startswith('## '):
            add_heading(doc, line[3:].strip(), 2)
        elif line.startswith('### '):
            add_heading(doc, line[4:].strip(), 3)
        elif line.startswith('- ') or line.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            text = line[2:].strip()
            # Simple bold parsing
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    p.add_run(part[2:-2]).bold = True
                else:
                    p.add_run(part)
        elif re.match(r'^\d+\.\s', line):
            p = doc.add_paragraph(style='List Number')
            text = re.sub(r'^\d+\.\s+', '', line)
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    p.add_run(part[2:-2]).bold = True
                else:
                    p.add_run(part)
        else:
            p = doc.add_paragraph()
            parts = re.split(r'(\*\*.*?\*\*)', line)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    p.add_run(part[2:-2]).bold = True
                else:
                    p.add_run(part)

    # 2. Add Lampiran (Appendix)
    doc.add_page_break()
    add_heading(doc, "LAMPIRAN", 1)
    
    p = doc.add_paragraph()
    p.add_run("Tabel 1. Rincian Agregat Temuan per Rumah Sakit").bold = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Read all JSON files
    data_dir = 'exports/agent_outputs/data_review/'
    files = [f for f in os.listdir(data_dir) if f.endswith('.json')]
    
    rs_data = []
    
    # Load Cochran to get RS Names
    import json
    master_path = 'data/Master_RS_Cochran_114.json'
    cochran_master = {}
    if os.path.exists(master_path):
        with open(master_path, 'r', encoding='utf-8') as f:
            cochran_master = json.load(f)

    for file in files:
        kode_rs = file.replace('data_review_', '').replace('.json', '')
        with open(os.path.join(data_dir, file), 'r', encoding='utf-8') as f:
            data = json.load(f)
            sampel = data.get('total_sample', 0)
            alerts = data.get('total_knavp_alerts', 0)
            sev = data.get('severity_counts', {})
            high = sev.get('High', 0)
            med = sev.get('Medium', 0)
            
            dual = 0
            for c in data.get('cases', []):
                if c.get('jumlah_beda_dual_coding', 0) > 0:
                    dual += 1
            
            nama_rs = cochran_master.get(kode_rs, {}).get('rs_name', f"RS {kode_rs}")
            
            rs_data.append({
                'kode': kode_rs,
                'nama': nama_rs,
                'sampel': sampel,
                'alerts': alerts,
                'high': high,
                'med': med,
                'dual': dual
            })
            
    # Create Table
    headers = ["No", "Kode RS", "Nama RS", "Total Sampel", "Alerts KNAVP", "Risiko Tinggi", "Risiko Sedang", "Diskrepansi Dual Coding"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        for paragraph in hdr_cells[i].paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
                run.font.size = Pt(10)

    # Add Data
    for idx, rs in enumerate(rs_data, 1):
        row_cells = table.add_row().cells
        row_cells[0].text = str(idx)
        row_cells[1].text = rs['kode']
        row_cells[2].text = rs['nama']
        row_cells[3].text = str(rs['sampel'])
        row_cells[4].text = str(rs['alerts'])
        row_cells[5].text = str(rs['high'])
        row_cells[6].text = str(rs['med'])
        row_cells[7].text = str(rs['dual'])
        
        for cell in row_cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(10)

    out_path = 'exports/Laporan_Akhir_Nasional_Agregat.docx'
    doc.save(out_path)
    print(f"DOCX created at {out_path}")

if __name__ == '__main__':
    main()
