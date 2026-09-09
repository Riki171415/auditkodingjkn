def update_export():
    with open('modules/export_generator.py', 'r', encoding='utf-8') as f:
        content = f.read()

    # 1. Replace E. Metode
    metode_bad = """    document.add_heading('E. Metode', level=2)
    add_p("Desk Review dilaksanakan melalui tahapan:", align=WD_ALIGN_PARAGRAPH.LEFT, size_pt=12, space_after_pt=4)
    metode_items = [
        ("1.", "ekstraksi data klaim;"),
        ("2.", "penerapan aturan validasi KNAVP;"),
        ("3.", "analisis hasil validasi oleh reviewer;"),
        ("4.", "dokumentasi hasil pada KKR-DR01; dan"),
        ("5.", "penyusunan rekomendasi tindak lanjut.")
    ]
    for num_str, text_str in metode_items:
        add_list_item(num_str, text_str)"""

    metode_good = """    document.add_heading('E. Metode', level=2)
    add_p("Pelaksanaan Desk Review dilakukan secara sistematis dan terstruktur yang dimulai dari ekstraksi data klaim hingga penyusunan rekomendasi tindak lanjut. Proses ini mencakup penerapan aturan validasi dari Katalog Nasional Aturan Validasi Pengodean (KNAVP), dilanjutkan dengan analisis mendalam terhadap hasil validasi oleh tim reviewer. Seluruh hasil temuan kemudian didokumentasikan ke dalam Kertas Kerja Review (KKR-DR01) sebagai dasar utama dalam penyusunan rekomendasi tindak lanjut bagi fasilitas kesehatan. Alur kerja tersebut dapat dilihat pada bagan berikut:", align=3, size_pt=12, space_after_pt=12)
    try:
        from docx.shared import Inches
        document.add_picture('assets/bagan_metode.png', width=Inches(6.0))
        last_paragraph = document.paragraphs[-1]
        last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_p("", space_after_pt=12)
    except Exception as e:
        print("Gagal menambahkan gambar bagan metode:", e)"""

    if metode_bad in content:
        content = content.replace(metode_bad, metode_good)
        print("E. Metode replaced!")
    else:
        print("E. Metode NOT FOUND!")
        
    # 2. Replace BAB III Kesimpulan
    bab3_bad = """    if dom_kat == '-' or total_temuan_all == 0:
        add_p(f"Pelaksanaan Desk Review terhadap data klaim {rs_name} menyimpulkan bahwa dari {len(cases)} kasus sampel yang diperiksa, kualitas pengodean telah terbukti sangat baik dan mematuhi aturan standar tanpa ada deteksi ketidaksesuaian berarti.", align=3, size_pt=12, space_after_pt=8)
    else:
        add_p(
            f"Pelaksanaan Desk Review terhadap data klaim {rs_name} menyimpulkan bahwa dari {len(cases)} kasus sampel yang diperiksa, terdapat indikasi ketidaksesuaian pengodean yang memerlukan perhatian khusus, terutama pada kelompok aturan {dom_label}. "
            "Sebagian kasus memiliki tingkat risiko tinggi terhadap kewajaran tarif dan pengelompokan iDRG, sehingga validasi lanjutan mutlak diperlukan untuk memastikan kesesuaian antara klaim dengan dokumen rekam medis pasien.",
            align=3, size_pt=12, space_after_pt=8
        )

    document.add_heading('B. Rekomendasi', level=2)
    add_p("Berdasarkan hasil analisis reviewer atas keseluruhan temuan Desk Review, direkomendasikan langkah-langkah tindak lanjut sebagai berikut:", align=WD_ALIGN_PARAGRAPH.LEFT, size_pt=12, space_after_pt=4)
    rekom_items = [
        ("●", "tidak diperlukan tindak lanjut bagi kasus yang lolos validasi tanpa catatan;"),
        ("●", "klarifikasi kepada PMIK terkait ketepatan penerapan aturan kombinasi kode dan pedoman ICS;"),
        ("●", "klarifikasi kepada DPJP terkait kelengkapan bukti medis penunjang diagnosis secondary/komplikasi;"),
        ("●", "pelaksanaan On-Site Audit bagi kasus-kasus prioritas dengan tingkat risiko tinggi dan temuan berulang;"),
        ("●", "pembinaan pengodean dan peningkatan pemahaman transisi iDRG secara berkala bagi tenaga koding rumah sakit.")
    ]
    for num_str, text_str in rekom_items:
        add_list_item(num_str, text_str)"""

    bab3_good = """    if dom_kat == '-' or total_temuan_all == 0:
        add_p(f"Pelaksanaan Desk Review terhadap data klaim {rs_name} menyimpulkan bahwa dari {len(cases)} kasus sampel yang diperiksa, kualitas pengodean telah terbukti sangat baik dan mematuhi aturan standar tanpa ada deteksi ketidaksesuaian berarti.", align=3, size_pt=12, space_after_pt=8)
    else:
        add_p(
            f"Pelaksanaan Desk Review terhadap data klaim {rs_name} menyimpulkan bahwa dari {len(cases)} kasus sampel yang diperiksa, terdapat indikasi ketidaksesuaian pengodean yang memerlukan perhatian khusus, terutama pada kelompok aturan {dom_label}. "
            f"{'Sebagian kasus memiliki tingkat risiko tinggi terhadap kewajaran tarif dan pengelompokan iDRG, sehingga validasi lanjutan mutlak diperlukan untuk memastikan kesesuaian antara klaim dengan dokumen rekam medis pasien.' if onsite > 0 else 'Mayoritas kasus dikategorikan ke dalam risiko rendah atau lolos validasi, namun perbaikan dan pembinaan koding tetap diperlukan.'}",
            align=3, size_pt=12, space_after_pt=8
        )

    document.add_heading('B. Rekomendasi', level=2)
    add_p("Berdasarkan hasil analisis reviewer atas keseluruhan temuan Desk Review, direkomendasikan langkah-langkah tindak lanjut sebagai berikut:", align=WD_ALIGN_PARAGRAPH.LEFT, size_pt=12, space_after_pt=4)
    rekom_items = [
        ("●", "tidak diperlukan tindak lanjut bagi kasus yang lolos validasi tanpa catatan;")
    ]
    if rule_counts.get('mutually_exclusive', 0) > 0 or rule_counts.get('underlying_manifestation', 0) > 0 or rule_counts.get('lainnya', 0) > 0:
        rekom_items.append(("●", "klarifikasi kepada PMIK terkait ketepatan penerapan aturan koding (termasuk kombinasi kode, excludes) dan kepatuhan terhadap pedoman ICS;"))
    if rule_counts.get('medical_evidence', 0) > 0 or rule_counts.get('procedure_validation', 0) > 0:
        rekom_items.append(("●", "klarifikasi kepada DPJP terkait kelengkapan bukti medis penunjang diagnosis secondary/komplikasi serta validitas tindakan;"))
    if onsite > 0:
        rekom_items.append(("●", f"pelaksanaan On-Site Audit terhadap {onsite} kasus prioritas dengan tingkat risiko tinggi dan/atau temuan berulang untuk memverifikasi dokumen rekam medis fisik;"))
    if sampling > 0:
        rekom_items.append(("●", f"pelaksanaan Klarifikasi/Sampling terhadap {sampling} kasus dengan risiko sedang untuk memastikan kelengkapan administratif;"))
        
    rekom_items.append(("●", "pembinaan pengodean dan peningkatan pemahaman transisi iDRG secara berkala bagi tenaga koding rumah sakit."))

    for num_str, text_str in rekom_items:
        add_list_item(num_str, text_str)"""

    if bab3_bad in content:
        content = content.replace(bab3_bad, bab3_good)
        print("BAB 3 replaced!")
    else:
        print("BAB 3 NOT FOUND!")

    with open('modules/export_generator.py', 'w', encoding='utf-8') as f:
        f.write(content)

if __name__ == '__main__':
    update_export()
