import openpyxl
from openpyxl.chart import BarChart, PieChart, Reference
import json
from modules.db_manager import get_recap_desk_review
import os

def create_excel_dashboard():
    print("Membuat Dashboard Grafik Nasional (Excel)...")
    recap = get_recap_desk_review()
    
    kep_counts = {'Direkomendasikan On-Site Audit': 0, 'Audit Sampling': 0, 'Lolos/Monitoring': 0}
    rule_counts = {}
    for r in recap:
        try: fd = json.loads(r.get('tindakan_reviewer') or '{}')
        except: fd = {}
        kep = fd.get('keputusan_sistem', 'Lolos/Monitoring')
        if kep in kep_counts: kep_counts[kep] += 1
        else: kep_counts['Lolos/Monitoring'] += 1
        
        triggered = r.get('triggered_rules', [])
        for rule in triggered:
            cat = rule.get('kelompok_rule', 'Lainnya')
            rule_counts[cat] = rule_counts.get(cat, 0) + 1
            
    sorted_rules = sorted(rule_counts.items(), key=lambda x: x[1], reverse=True)[:5]
    if not sorted_rules: sorted_rules = [("Tidak ada", 1)]

    wb = openpyxl.Workbook()
    
    # Sheet 1: Dashboard Validasi (Bar Chart Kategori)
    ws1 = wb.active
    ws1.title = "Dashboard Kategori Rule"
    ws1.append(["Kategori Rule", "Jumlah Kasus"])
    for k, v in sorted_rules:
        ws1.append([k, v])
        
    chart1 = BarChart()
    chart1.type = "col"
    chart1.title = "Top 5 Kategori Rule Terpicu"
    chart1.y_axis.title = "Jumlah Kasus"
    chart1.x_axis.title = "Kategori"
    
    data1 = Reference(ws1, min_col=2, min_row=1, max_row=len(sorted_rules)+1)
    cats1 = Reference(ws1, min_col=1, min_row=2, max_row=len(sorted_rules)+1)
    chart1.add_data(data1, titles_from_data=True)
    chart1.set_categories(cats1)
    ws1.add_chart(chart1, "D2")
    
    # Sheet 2: Dashboard Keputusan (Pie Chart)
    ws2 = wb.create_sheet(title="Dashboard Keputusan")
    ws2.append(["Keputusan", "Jumlah Kasus"])
    for k, v in kep_counts.items():
        ws2.append([k, v])
        
    chart2 = PieChart()
    chart2.title = "Distribusi Keputusan KNAVP Nasional"
    data2 = Reference(ws2, min_col=2, min_row=1, max_row=len(kep_counts)+1)
    cats2 = Reference(ws2, min_col=1, min_row=2, max_row=len(kep_counts)+1)
    chart2.add_data(data2, titles_from_data=True)
    chart2.set_categories(cats2)
    ws2.add_chart(chart2, "D2")
    
    os.makedirs('exports/word_reports', exist_ok=True)
    out_path = 'exports/word_reports/Dashboard_Grafik_Nasional.xlsx'
    wb.save(out_path)
    print(f"Berhasil membuat {out_path}")

if __name__ == '__main__':
    create_excel_dashboard()
