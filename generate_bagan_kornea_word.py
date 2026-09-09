import os
import matplotlib.pyplot as plt
import matplotlib.patches as patches
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUTPUT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'exports', 'kornea_docs')
os.makedirs(OUTPUT_DIR, exist_ok=True)

def draw_box(ax, x, y, width, height, text, facecolor='#f8fafc', edgecolor='#334155', fontsize=8.5, weight='bold', text_color='#1e293b'):
    """Helper to draw rounded rectangle box with centered text."""
    box = patches.FancyBboxPatch(
        (x - width/2, y - height/2),
        width, height,
        boxstyle="round,pad=0.2,rounding_size=0.15",
        fc=facecolor, ec=edgecolor, lw=1.5, zorder=3
    )
    ax.add_patch(box)
    ax.text(x, y, text, ha='center', va='center', fontsize=fontsize, fontweight=weight, color=text_color, zorder=4, linespacing=1.2)

def draw_arrow(ax, x1, y1, x2, y2, color='#475569', lw=1.5, style='->'):
    """Helper to draw arrow from (x1, y1) to (x2, y2)."""
    ax.annotate('', xy=(x2, y2), xytext=(x1, y1),
                arrowprops=dict(arrowstyle=style, color=color, lw=lw, shrinkA=0, shrinkB=0),
                zorder=2)

def generate_bagan1_png():
    print("Generating Bagan 1 PNG...")
    fig, ax = plt.subplots(figsize=(10, 15), dpi=300)
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 15)
    ax.axis('off')

    # Main linear flow nodes (top down)
    # Y coordinates: 14.2, 13.1, 12.0, 10.9, 9.8, 8.5, 7.1, 5.8
    nodes = [
        (14.2, "INFORMASI DARI KELUARGA /\nKOORDINATOR DONOR", '#f1f5f9', '#475569'),
        (13.1, "VERIFIKASI DOKUMEN DONOR\n(Authorization, Medical Record, Riwayat Donor)", '#f1f5f9', '#475569'),
        (12.0, "PERSIAPAN PERLENGKAPAN", '#f1f5f9', '#475569'),
        (10.9, "TIM BANK MATA MENUJU LOKASI DONOR", '#f1f5f9', '#475569'),
        (9.8,  "INFORMED CONSENT / LEGAL AUTHORIZATION", '#f1f5f9', '#475569'),
        (8.5,  "PEMERIKSAAN FISIK DONOR DAN PEMERIKSAAN AWAL\nKORNEA (PENLIGHT)\n(Kejernihan kornea, defek epitel, benda asing, kontaminasi, warna sklera)", '#dcfce3', '#16a34a'),
        (7.1,  "MELAKUKAN TINDAKAN ASEPTIK PADA AREA MATA DONOR\n(Pembersihan kelopak mata, bulu mata, dan area periokular)", '#dcfce3', '#16a34a'),
        (5.8,  "RECOVERY / PENGAMBILAN JARINGAN KORNEA", '#dcfce3', '#16a34a'),
    ]

    for y, text, fc, ec in nodes:
        h = 0.9 if '\n\n' in text or text.count('\n') >= 2 else 0.7
        draw_box(ax, 5.0, y, 6.5, h, text, facecolor=fc, edgecolor=ec, fontsize=8.5)

    for i in range(len(nodes) - 1):
        y1 = nodes[i][0] - (0.45 if nodes[i][0] in [8.5, 7.1] else 0.35)
        y2 = nodes[i+1][0] + (0.45 if nodes[i+1][0] in [8.5, 7.1] else 0.35)
        draw_arrow(ax, 5.0, y1, 5.0, y2)

    # Branching from Recovery (y=5.8)
    y_rec_bottom = 5.8 - 0.35
    # Left Branch (Kornea / Optisol)
    draw_arrow(ax, 5.0, y_rec_bottom, 2.7, 4.6)
    draw_box(ax, 2.7, 4.2, 4.2, 0.8, "KORNEA DISIMPAN DALAM\nOPTISOL GS (Suhu 2 – 8 °C)", facecolor='#f3e8ff', edgecolor='#9333ea', fontsize=8.5)
    draw_arrow(ax, 2.7, 3.8, 2.7, 2.8)
    draw_box(ax, 2.7, 2.4, 4.2, 0.8, "KORNEA DISIMPAN DI\nCOOL BOX (Suhu 2 – 8 °C)", facecolor='#f3e8ff', edgecolor='#9333ea', fontsize=8.5)

    # Right Branch (Darah / Vacutainer)
    draw_arrow(ax, 5.0, y_rec_bottom, 7.3, 4.6)
    draw_box(ax, 7.3, 4.2, 4.2, 0.8, "PENGAMBILAN SAMPEL DARAH DONOR\n(Untuk pemeriksaan penyakit infeksi sesuai standar)", facecolor='#ffedd5', edgecolor='#ea580c', fontsize=8.5)
    draw_arrow(ax, 7.3, 3.8, 7.3, 2.8)
    draw_box(ax, 7.3, 2.4, 4.2, 0.8, "DARAH DISIMPAN DALAM\nTABUNG VACUTAINER (Suhu 2 – 8 °C)", facecolor='#ffedd5', edgecolor='#ea580c', fontsize=8.5)
    draw_arrow(ax, 7.3, 2.0, 7.3, 1.0)
    draw_box(ax, 7.3, 0.6, 4.2, 0.8, "DARAH DISIMPAN DI\nCOOL BOX (Suhu 2 – 8 °C)", facecolor='#ffedd5', edgecolor='#ea580c', fontsize=8.5)

    ax.set_title("BAGAN 1: ALUR RECOVERY DAN PRESERVASI JARINGAN KORNEA", fontsize=13, fontweight='bold', pad=20, color='#0f172a')
    plt.tight_layout()
    out_path = os.path.join(OUTPUT_DIR, 'bagan1_kornea.png')
    plt.savefig(out_path, dpi=300, bbox_inches='tight', facecolor='white')
    plt.close()
    return out_path

def generate_bagan2_png():
    print("Generating Bagan 2 PNG...")
    fig, ax = plt.subplots(figsize=(14, 18), dpi=300)
    ax.set_xlim(0, 14)
    ax.set_ylim(0, 18)
    ax.axis('off')

    # Start Node
    draw_box(ax, 7.0, 17.2, 4.5, 0.7, "KORNEA DONOR", facecolor='#e2e8f0', edgecolor='#334155', fontsize=10)

    # Branch Left (Fresh Donor x=3.5) and Right (Kiriman Pusat x=10.5)
    draw_arrow(ax, 7.0, 16.85, 3.5, 16.1)
    draw_arrow(ax, 7.0, 16.85, 10.5, 16.1)

    # --- LEFT BRANCH: FRESH DONOR ---
    draw_box(ax, 3.5, 15.7, 4.2, 0.7, "FRESH DONOR", facecolor='#dbeafe', edgecolor='#2563eb', fontsize=9.5)
    draw_arrow(ax, 3.5, 15.35, 3.5, 14.65)
    draw_box(ax, 3.5, 14.3, 4.6, 0.7, "VERIFIKASI DOKUMEN DONOR", facecolor='#f1f5f9', edgecolor='#475569', fontsize=8.5)
    draw_arrow(ax, 3.5, 13.95, 3.5, 13.25)
    draw_box(ax, 3.5, 12.9, 4.6, 0.7, "MEMERIKSA KEADAAN BOTOL KORNEA", facecolor='#f1f5f9', edgecolor='#475569', fontsize=8.5)

    # Branch from Botol Fresh: Damaged vs Good
    draw_arrow(ax, 3.5, 12.55, 1.4, 11.65) # Damaged left
    draw_arrow(ax, 3.5, 12.55, 4.8, 11.65) # Good right

    draw_box(ax, 1.4, 11.2, 2.6, 0.9, "BOTOL RETAK/PECAH\nATAU RUSAK", facecolor='#fee2e2', edgecolor='#dc2626', fontsize=8)
    draw_box(ax, 4.8, 11.2, 3.2, 0.9, "TIDAK TERDAPAT\nKERUSAKAN", facecolor='#dcfce3', edgecolor='#16a34a', fontsize=8)

    # Damaged flow 1
    draw_arrow(ax, 1.4, 10.75, 1.4, 9.95)
    draw_box(ax, 1.4, 9.6, 2.2, 0.7, "REJECTED", facecolor='#ef4444', edgecolor='#b91c1c', fontsize=8.5, text_color='white')
    draw_arrow(ax, 1.4, 9.25, 0.8, 8.45)
    draw_arrow(ax, 1.4, 9.25, 2.0, 8.45)
    draw_box(ax, 0.8, 8.1, 1.4, 0.7, "INSTALASI\nKESLING", facecolor='#fef2f2', edgecolor='#f87171', fontsize=7.5)
    draw_box(ax, 2.0, 8.1, 1.4, 0.7, "WET LAB", facecolor='#fef2f2', edgecolor='#f87171', fontsize=7.5)

    # Good flow -> Cek Darah
    draw_arrow(ax, 4.8, 10.75, 4.8, 9.95)
    draw_box(ax, 4.8, 9.5, 3.6, 0.9, "PEMERIKSAAN SAMPLE\nDARAH DONOR (LAB)", facecolor='#fef3c7', edgecolor='#d97706', fontsize=8.5)

    # Cek Darah -> Positif vs Negatif
    draw_arrow(ax, 4.8, 9.05, 3.2, 8.25) # Positif
    draw_arrow(ax, 4.8, 9.05, 5.8, 8.25) # Negatif

    draw_box(ax, 3.2, 7.9, 1.8, 0.7, "POSITIF (+)", facecolor='#fee2e2', edgecolor='#dc2626', fontsize=8)
    draw_arrow(ax, 3.2, 7.55, 3.2, 6.85)
    draw_box(ax, 3.2, 6.5, 2.0, 0.7, "REJECTED", facecolor='#ef4444', edgecolor='#b91c1c', fontsize=8.5, text_color='white')
    draw_arrow(ax, 3.2, 6.15, 2.5, 5.45)
    draw_arrow(ax, 3.2, 6.15, 3.9, 5.45)
    draw_box(ax, 2.5, 5.1, 1.3, 0.7, "INSTALASI\nKESLING", facecolor='#fef2f2', edgecolor='#f87171', fontsize=7.5)
    draw_box(ax, 3.9, 5.1, 1.3, 0.7, "WET LAB", facecolor='#fef2f2', edgecolor='#f87171', fontsize=7.5)

    draw_box(ax, 5.8, 7.9, 1.8, 0.7, "NEGATIF (-)", facecolor='#dcfce3', edgecolor='#16a34a', fontsize=8)
    draw_arrow(ax, 5.8, 7.55, 5.8, 6.85)
    draw_box(ax, 5.8, 6.5, 3.4, 0.7, "PEMERIKSAAN DENGAN\nSLIT LAMP", facecolor='#e0e7ff', edgecolor='#4f46e5', fontsize=8.5)
    draw_arrow(ax, 5.8, 6.15, 5.8, 5.45)
    draw_box(ax, 5.8, 5.1, 3.4, 0.7, "PEMERIKSAAN SPECULAR\nMIKROSKOP", facecolor='#e0e7ff', edgecolor='#4f46e5', fontsize=8.5)
    draw_arrow(ax, 5.8, 4.75, 5.8, 4.05)
    draw_box(ax, 5.8, 3.6, 3.8, 0.9, "DISIMPAN DI REFRIGERATOR\nUNIT BANK MATA", facecolor='#dcfce3', edgecolor='#15803d', fontsize=8.5)


    # --- RIGHT BRANCH: KIRIMAN PUSAT ---
    draw_box(ax, 10.5, 15.7, 4.8, 0.7, "KIRIMAN DARI BANK MATA PUSAT", facecolor='#f3e8ff', edgecolor='#9333ea', fontsize=9.5)
    draw_arrow(ax, 10.5, 15.35, 10.5, 14.65)
    draw_box(ax, 10.5, 14.3, 4.6, 0.7, "VERIFIKASI DOKUMEN DONOR", facecolor='#f1f5f9', edgecolor='#475569', fontsize=8.5)
    draw_arrow(ax, 10.5, 13.95, 10.5, 13.25)
    draw_box(ax, 10.5, 12.9, 4.6, 0.7, "MEMERIKSA KEADAAN BOTOL KORNEA", facecolor='#f1f5f9', edgecolor='#475569', fontsize=8.5)

    # Branch from Botol Pusat: Damaged vs Good
    draw_arrow(ax, 10.5, 12.55, 8.8, 11.65) # Damaged left
    draw_arrow(ax, 10.5, 12.55, 12.2, 11.65) # Good right

    draw_box(ax, 8.8, 11.2, 2.6, 0.9, "BOTOL RETAK/PECAH\nATAU RUSAK", facecolor='#fee2e2', edgecolor='#dc2626', fontsize=8)
    draw_box(ax, 12.2, 11.2, 3.2, 0.9, "TIDAK TERDAPAT\nKERUSAKAN", facecolor='#dcfce3', edgecolor='#16a34a', fontsize=8)

    # Damaged flow 2
    draw_arrow(ax, 8.8, 10.75, 8.8, 9.95)
    draw_box(ax, 8.8, 9.6, 2.2, 0.7, "REJECTED", facecolor='#ef4444', edgecolor='#b91c1c', fontsize=8.5, text_color='white')
    draw_arrow(ax, 8.8, 9.25, 8.1, 8.45)
    draw_arrow(ax, 8.8, 9.25, 9.5, 8.45)
    draw_box(ax, 8.1, 8.1, 1.4, 0.7, "INSTALASI\nKESLING", facecolor='#fef2f2', edgecolor='#f87171', fontsize=7.5)
    draw_box(ax, 9.5, 8.1, 1.4, 0.7, "WET LAB", facecolor='#fef2f2', edgecolor='#f87171', fontsize=7.5)

    # Good flow Pusat -> Simpan
    draw_arrow(ax, 12.2, 10.75, 12.2, 4.25)
    draw_box(ax, 12.2, 3.6, 3.8, 0.9, "DISIMPAN DI REFRIGERATOR\nUNIT BANK MATA", facecolor='#dcfce3', edgecolor='#15803d', fontsize=8.5)


    # --- MERGING TO RESIPIEN & TRANSPLANTASI ---
    draw_arrow(ax, 5.8, 3.15, 8.0, 2.35)
    draw_arrow(ax, 12.2, 3.15, 9.6, 2.35)

    draw_box(ax, 8.8, 2.0, 3.6, 0.7, "RESIPIEN", facecolor='#ffedd5', edgecolor='#ea580c', fontsize=9.5)
    draw_arrow(ax, 8.8, 1.65, 8.8, 0.95)
    draw_box(ax, 8.8, 0.6, 4.0, 0.7, "TRANSPLANTASI", facecolor='#d1fae5', edgecolor='#059669', fontsize=10.5)

    ax.set_title("BAGAN 2: ALUR PENYIMPANAN JARINGAN KORNEA", fontsize=14, fontweight='bold', pad=20, color='#0f172a')
    plt.tight_layout()
    out_path = os.path.join(OUTPUT_DIR, 'bagan2_kornea.png')
    plt.savefig(out_path, dpi=300, bbox_inches='tight', facecolor='white')
    plt.close()
    return out_path

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def build_word_document(bagan1_png, bagan2_png):
    print("Building MS Word Document (.docx)...")
    doc = Document()

    # Set margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Header title
    p_hdr = doc.add_paragraph()
    p_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_hdr = p_hdr.add_run("STANDARD OPERATING PROCEDURE & FLOWCHART\nALUR RECOVERY DAN PENYIMPANAN JARINGAN KORNEA")
    run_hdr.bold = True
    run_hdr.font.size = Pt(14)
    run_hdr.font.color.rgb = RGBColor(15, 23, 42)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run("BANK MATA INDONESIA / PELAYANAN TRANSPLANTASI KORNEA\n(Dilengkapi Diagram Visual & Teks Siap Copy-Paste untuk SmartArt / Shapes MS Word)")
    run_sub.italic = True
    run_sub.font.size = Pt(10)
    run_sub.font.color.rgb = RGBColor(71, 85, 105)

    doc.add_paragraph().add_run("---")

    # =========================================================
    # BAGAN 1 SECTION
    # =========================================================
    h1 = doc.add_heading("BAGAN 1: ALUR RECOVERY DAN PRESERVASI JARINGAN KORNEA", level=1)
    h1.runs[0].font.color.rgb = RGBColor(30, 58, 138)

    p_intro1 = doc.add_paragraph("Berikut adalah visualisasi diagram alur (Flowchart) dan teks terstruktur untuk prosedur pemulihan (recovery) dan preservasi jaringan kornea dari donor. Anda dapat langsung melihat diagram di bawah ini atau menyalin teks terstruktur untuk dimasukkan ke dalam SmartArt / Shapes MS Word Anda.")
    
    # Insert diagram image 1
    p_img1 = doc.add_paragraph()
    p_img1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img1.add_run().add_picture(bagan1_png, width=Inches(5.5))
    p_cap1 = doc.add_paragraph()
    p_cap1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_cap1 = p_cap1.add_run("Gambar 1: Diagram Alur Recovery dan Preservasi Jaringan Kornea (Resolusi Tinggi)")
    run_cap1.italic = True
    run_cap1.font.size = Pt(9)

    doc.add_heading("Teks Alur untuk Copy-Paste ke MS Word (SmartArt / Shapes)", level=2)
    doc.add_paragraph("Gunakan teks terstruktur di bawah ini jika Anda ingin membuat bagan secara manual di Word agar 100% bisa diedit langsung pada bentuk (shapes) atau SmartArt:")

    table1 = doc.add_table(rows=1, cols=1)
    table1.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell1 = table1.cell(0, 0)
    set_cell_background(cell1, 'F8FAFC')
    cell1.width = Inches(6.8)

    text_bagan1 = """1. INFORMASI DARI KELUARGA / KOORDINATOR DONOR
↓
2. VERIFIKASI DOKUMEN DONOR (Authorization, Medical Record, Riwayat Donor)
↓
3. PERSIAPAN PERLENGKAPAN
↓
4. TIM BANK MATA MENUJU LOKASI DONOR
↓
5. INFORMED CONSENT / LEGAL AUTHORIZATION
↓
6. PEMERIKSAAN FISIK DONOR DAN PEMERIKSAAN AWAL KORNEA (PENLIGHT)
   (Untuk menilai kondisi donor dan kornea, kejernihan kornea, defek epitel, benda asing, kontaminasi, dan warna sklera)
↓
7. MELAKUKAN TINDAKAN ASEPTIK PADA AREA MATA DONOR
   (Pembersihan kelopak mata, bulu mata, dan area periokular sesuai prosedur aseptik)
↓
8. RECOVERY / PENGAMBILAN JARINGAN KORNEA
↓
[Dari langkah nomor 8 bercabang menjadi 2 (dua) arah paralel ke bawah]:

► CABANG KIRI (PRESERVASI KORNEA):
   9A. KORNEA DISIMPAN DALAM OPTISOL GS (Suhu 2 – 8 °C)
   ↓
   10A. KORNEA DISIMPAN DI COOL BOX (Suhu 2 – 8 °C)

► CABANG KANAN (PRESERVASI SAMPEL DARAH):
   9B. PENGAMBILAN SAMPEL DARAH DONOR (Untuk pemeriksaan penyakit infeksi sesuai standar unit bank mata)
   ↓
   10B. DARAH DISIMPAN DALAM TABUNG VACUTAINER (Suhu 2 – 8 °C)
   ↓
   11B. DARAH DISIMPAN DI COOL BOX (Suhu 2 – 8 °C)"""

    p_t1 = cell1.paragraphs[0]
    p_t1.paragraph_format.left_indent = Inches(0.1)
    p_t1.paragraph_format.right_indent = Inches(0.1)
    run_t1 = p_t1.add_run(text_bagan1)
    run_t1.font.name = 'Consolas'
    run_t1.font.size = Pt(9.5)
    run_t1.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_page_break()

    # =========================================================
    # BAGAN 2 SECTION
    # =========================================================
    h2 = doc.add_heading("BAGAN 2: ALUR PENYIMPANAN JARINGAN KORNEA", level=1)
    h2.runs[0].font.color.rgb = RGBColor(30, 58, 138)

    p_intro2 = doc.add_paragraph("Diagram dan alur berikut memetakan proses penerimaan, verifikasi kondisi botol, pemeriksaan serologi darah donor, pemeriksaan mikroskopis kornea, hingga penyimpanan jaringan di refrigerator sebelum dilanjutkan ke tahap transplantasi pada resipien.")

    # Insert diagram image 2
    p_img2 = doc.add_paragraph()
    p_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img2.add_run().add_picture(bagan2_png, width=Inches(6.3))
    p_cap2 = doc.add_paragraph()
    p_cap2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_cap2 = p_cap2.add_run("Gambar 2: Diagram Alur Penyimpanan Jaringan Kornea (Resolusi Tinggi)")
    run_cap2.italic = True
    run_cap2.font.size = Pt(9)

    doc.add_heading("Teks Alur untuk Copy-Paste ke MS Word (SmartArt / Shapes)", level=2)
    doc.add_paragraph("Salin teks terstruktur di bawah ini untuk membuat bagan penyimpanan jaringan kornea pada shapes / SmartArt MS Word:")

    table2 = doc.add_table(rows=1, cols=1)
    table2.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell2 = table2.cell(0, 0)
    set_cell_background(cell2, 'F8FAFC')
    cell2.width = Inches(6.8)

    text_bagan2 = """1. KORNEA DONOR
↓
[Dari nomor 1 bercabang menjadi 2 (dua) arah ke bawah: FRESH DONOR dan KIRIMAN PUSAT]

====================================================================
► CABANG KIRI: FRESH DONOR
====================================================================
2A. FRESH DONOR
↓
3A. VERIFIKASI DOKUMEN DONOR
↓
4A. MEMERIKSA KEADAAN BOTOL KORNEA
↓
[Percabangan berdasarkan kondisi botol]:
   ├─► BOTOL RETAK / PECAH ATAU TERDAPAT KERUSAKAN
   │   ↓
   │   REJECTED -> (Bercabang ke) INSTALASI KESLING dan WET LAB
   │
   └─► TIDAK TERDAPAT KERUSAKAN
       ↓
       5A. PEMERIKSAAN SAMPLE DARAH DONOR (LABORATORIUM)
       ↓
       [Percabangan berdasarkan hasil pemeriksaan darah]:
          ├─► POSITIF (+)
          │   ↓
          │   REJECTED -> (Bercabang ke) INSTALASI KESLING dan WET LAB
          │
          └─► NEGATIF (-)
              ↓
              6A. PEMERIKSAAN DENGAN SLIT LAMP
              ↓
              7A. PEMERIKSAAN SPECULAR MIKROSKOP
              ↓
              8A. DISIMPAN DI REFRIGERATOR UNIT BANK MATA

====================================================================
► CABANG KANAN: KIRIMAN DARI BANK MATA PUSAT
====================================================================
2B. KIRIMAN DARI BANK MATA PUSAT
↓
3B. VERIFIKASI DOKUMEN DONOR
↓
4B. MEMERIKSA KEADAAN BOTOL KORNEA
↓
[Percabangan berdasarkan kondisi botol]:
   ├─► BOTOL RETAK / PECAH ATAU TERDAPAT KERUSAKAN
   │   ↓
   │   REJECTED -> (Bercabang ke) INSTALASI KESLING dan WET LAB
   │
   └─► TIDAK TERDAPAT KERUSAKAN
       ↓
       5B. DISIMPAN DI REFRIGERATOR UNIT BANK MATA

====================================================================
► PENGGABUNGAN (MERGING MENUJU RESIPIEN & TRANSPLANTASI)
====================================================================
Dari kedua kotak akhir penyimpanan:
( 8A. DISIMPAN DI REFRIGERATOR UNIT BANK MATA ) ──┐
                                                 ├─► RESIPIEN ──► TRANSPLANTASI
( 5B. DISIMPAN DI REFRIGERATOR UNIT BANK MATA ) ──┘"""

    p_t2 = cell2.paragraphs[0]
    p_t2.paragraph_format.left_indent = Inches(0.1)
    p_t2.paragraph_format.right_indent = Inches(0.1)
    run_t2 = p_t2.add_run(text_bagan2)
    run_t2.font.name = 'Consolas'
    run_t2.font.size = Pt(9.0)
    run_t2.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_page_break()

    # =========================================================
    # LAMPIRAN: KODE MERMAID
    # =========================================================
    h3 = doc.add_heading("LAMPIRAN: KODE DIAGRAM MERMAID (OPTIONAL GENERATOR)", level=1)
    h3.runs[0].font.color.rgb = RGBColor(30, 58, 138)
    doc.add_paragraph("Anda juga dapat menyalin kode Mermaid di bawah ini ke web diagram generator (seperti mermaid.live atau draw.io) untuk membuat variasi warna atau bentuk diagram lainnya.")

    doc.add_heading("1. Kode Mermaid - Bagan 1: Alur Recovery", level=2)
    table_m1 = doc.add_table(rows=1, cols=1)
    table_m1.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_m1 = table_m1.cell(0, 0)
    set_cell_background(cell_m1, 'F1F5F9')
    cell_m1.width = Inches(6.8)
    
    code_m1 = """graph TD
    A[INFORMASI DARI KELUARGA / KOORDINATOR DONOR] --> B[VERIFIKASI DOKUMEN DONOR<br/>Authorization, Medical Record, Riwayat Donor]
    B --> C[PERSIAPAN PERLENGKAPAN]
    C --> D[TIM BANK MATA MENUJU LOKASI DONOR]
    D --> E[INFORMED CONSENT / LEGAL AUTHORIZATION]
    
    E --> F[PEMERIKSAAN FISIK DONOR DAN<br/>PEMERIKSAAN AWAL KORNEA PENLIGHT<br/>Untuk menilai kondisi donor dan kornea,<br/>kejernihan kornea, defek epitel, benda asing,<br/>kontaminasi, dan warna sklera]
    style F fill:#dcfce3,stroke:#86efac,stroke-width:1px
    
    F --> G[MELAKUKAN TINDAKAN ASEPTIK<br/>PADA AREA MATA DONOR<br/>Pembersihan kelopak mata, bulu mata,<br/>dan area periokular sesuai prosedur aseptik]
    style G fill:#dcfce3,stroke:#86efac,stroke-width:1px
    
    G --> H[RECOVERY / PENGAMBILAN<br/>JARINGAN KORNEA]
    style H fill:#dcfce3,stroke:#86efac,stroke-width:1px
    
    H --> I[KORNEA DISIMPAN DALAM OPTISOL GS<br/>Suhu 2 – 8 °C]
    style I fill:#f3e8ff,stroke:#d8b4fe,stroke-width:1px
    
    I --> J[KORNEA DISIMPAN DI<br/>COOL BOX<br/>Suhu 2 – 8 °C]
    style J fill:#f3e8ff,stroke:#d8b4fe,stroke-width:1px
    
    H --> K[PENGAMBILAN SAMPEL DARAH DONOR<br/>Untuk pemeriksaan penyakit infeksi<br/>sesuai standar unit bank mata]
    style K fill:#ffedd5,stroke:#fdba74,stroke-width:1px
    
    K --> L[DARAH DISIMPAN DALAM<br/>TABUNG VACUTAINER<br/>Suhu 2 – 8 °C]
    style L fill:#ffedd5,stroke:#fdba74,stroke-width:1px
    
    L --> M[DARAH DISIMPAN DI<br/>COOL BOX<br/>Suhu 2 – 8 °C]
    style M fill:#ffedd5,stroke:#fdba74,stroke-width:1px"""
    run_m1 = cell_m1.paragraphs[0].add_run(code_m1)
    run_m1.font.name = 'Consolas'
    run_m1.font.size = Pt(8.0)

    doc.add_heading("2. Kode Mermaid - Bagan 2: Alur Penyimpanan", level=2)
    table_m2 = doc.add_table(rows=1, cols=1)
    table_m2.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_m2 = table_m2.cell(0, 0)
    set_cell_background(cell_m2, 'F1F5F9')
    cell_m2.width = Inches(6.8)

    code_m2 = """graph TD
    Start[KORNEA DONOR] --> Fresh[FRESH DONOR]
    Start --> Pusat[KIRIMAN DARI BANK MATA PUSAT]

    %% Jalur Kiri (Fresh Donor)
    Fresh --> VerifikasiFresh[VERIFIKASI DOKUMEN DONOR]
    VerifikasiFresh --> CekBotolFresh[MEMERIKSA KEADAAN BOTOL KORNEA]
    
    CekBotolFresh --> BotolRusakFresh[BOTOL RETAK/ PECAH<br/>ATAU TERDAPAT<br/>KERUSAKAN]
    CekBotolFresh --> BotolBagusFresh[TIDAK TERDAPAT<br/>KERUSAKAN]

    BotolRusakFresh --> RejectedFresh1[REJECTED]
    RejectedFresh1 --> KeslingFresh1[INSTALASI<br/>KESLING]
    RejectedFresh1 --> WetlabFresh1[WET LAB]

    BotolBagusFresh --> CekDarah[PEMERIKSAAN SAMPLE<br/>DARAH DONOR<br/>LABORATORIUM]
    CekDarah --> Positif[POSITIF]
    CekDarah --> Negatif[NEGATIF]

    Positif --> RejectedFresh2[REJECTED]
    RejectedFresh2 --> KeslingFresh2[INSTALASI<br/>KESLING]
    RejectedFresh2 --> WetlabFresh2[WET LAB]

    Negatif --> SlitLamp[PEMERIKSAAN DENGAN<br/>SLIT LAMP]
    SlitLamp --> Specular[PEMERIKSAAN<br/>SPECULAR MIKROSKOP]
    Specular --> SimpanKulkasFresh[DISIMPAN DI REFRIGERATOR<br/>UNIT BANK MATA]

    %% Jalur Kanan (Kiriman Pusat)
    Pusat --> VerifikasiPusat[VERIFIKASI DOKUMEN DONOR]
    VerifikasiPusat --> CekBotolPusat[MEMERIKSA KEADAAN BOTOL KORNEA]

    CekBotolPusat --> BotolRusakPusat[BOTOL RETAK/ PECAH<br/>ATAU TERDAPAT<br/>KERUSAKAN]
    CekBotolPusat --> BotolBagusPusat[TIDAK TERDAPAT<br/>KERUSAKAN]

    BotolRusakPusat --> RejectedPusat[REJECTED]
    RejectedPusat --> KeslingPusat[INSTALASI<br/>KESLING]
    RejectedPusat --> WetlabPusat[WET LAB]

    BotolBagusPusat --> SimpanKulkasPusat[DISIMPAN DI REFRIGERATOR<br/>UNIT BANK MATA]

    %% Penggabungan ke Resipien
    SimpanKulkasFresh --> Resipien[RESIPIEN]
    SimpanKulkasPusat --> Resipien

    Resipien --> Transplantasi[TRANSPLANTASI]"""
    run_m2 = cell_m2.paragraphs[0].add_run(code_m2)
    run_m2.font.name = 'Consolas'
    run_m2.font.size = Pt(8.0)

    # Save Word document
    docx_path = os.path.join(OUTPUT_DIR, "Alur_Recovery_dan_Penyimpanan_Jaringan_Kornea.docx")
    doc.save(docx_path)
    print(f"SELESAI! Dokumen MS Word tersimpan di:\n{docx_path}")
    return docx_path

if __name__ == '__main__':
    b1_png = generate_bagan1_png()
    b2_png = generate_bagan2_png()
    docx_file = build_word_document(b1_png, b2_png)
