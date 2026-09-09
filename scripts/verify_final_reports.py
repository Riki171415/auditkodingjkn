"""Verify all Word, Excel, and snapshot consistency — filenames include RS name."""
import json
import re
import sys
import zipfile
from pathlib import Path
from collections import Counter

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from docx import Document
import openpyxl

OUT = Path(__file__).resolve().parents[1] / 'outputs' / 'laporan_final_20260831'
snap = json.loads((OUT / 'snapshot.json').read_text(encoding='utf-8'))


def rs_slug(nama_rs, max_len=40):
    slug = re.sub(r'[^A-Z0-9]+', '_', str(nama_rs).upper().strip())
    return slug.strip('_')[:max_len].rstrip('_')


def sheet_rows(path, sheet):
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    data = list(wb[sheet].values)
    wb.close()
    return data


METRICS = ['total', 'onsite', 'sampling', 'monitoring', 'unknown', 'dc_cases', 'dc_total', 'alerts', 'avg_score']
national = sheet_rows(OUT / 'Rekap_Nasional.xlsx', 'Ringkasan Eksekutif')
national_by_rs = {str(r[0]): r for r in national[1:]}
national_detail = sheet_rows(OUT / 'Rekap_Nasional.xlsx', 'Master Data (Rincian)')
hdr = national_detail[0]
national_cases = {
    (str(r[hdr.index('Kode RS Sumber')]), str(r[hdr.index('Nomor SEP')])): r[hdr.index('Rekomendasi Laporan')]
    for r in national_detail[1:]
}

observed = Counter()
errors = []

for h in snap['hospitals']:
    code = h['kode_rs']
    slug = rs_slug(h['nama_rs'])
    x = OUT / 'excel_per_rs' / f'Rekap_{code}_{slug}.xlsx'
    d = OUT / 'word_per_rs' / f'LHR_{code}_{slug}.docx'

    if not x.exists():
        errors.append(f'MISSING Excel {x.name}'); continue
    if not d.exists():
        errors.append(f'MISSING Word {d.name}'); continue

    # Excel per RS: ringkasan
    summary = sheet_rows(x, 'Ringkasan Eksekutif RS')[1]
    expected = [h['summary'][k] for k in METRICS]
    if list(summary[2:11]) != expected:
        errors.append(f'FAIL excel summary {code}: {list(summary[2:11])} != {expected}')
    if list(national_by_rs[code][2:11]) != expected:
        errors.append(f'FAIL national row {code}: {list(national_by_rs[code][2:11])} != {expected}')
    if summary[11] != 0:
        errors.append(f'FAIL selisih {code}: {summary[11]}')

    # Word per RS
    doc = Document(str(d))
    wc = len(doc.tables[1].rows) - 1
    if wc != h['summary']['total']:
        errors.append(f'FAIL word count {code}: {wc} != {h["summary"]["total"]}')

    word_cases = {r.cells[1].text: r.cells[12].text for r in doc.tables[1].rows[1:]}
    detail = sheet_rows(x, 'Rincian SEP Kasus'); heads = detail[0]
    for r in detail[1:]:
        sep = str(r[heads.index('Nomor SEP')]); decision = r[heads.index('Rekomendasi Laporan')]
        nat_dec = national_cases.get((code, sep)); word_dec = word_cases.get(sep)
        if decision != nat_dec or decision != word_dec:
            errors.append(f'FAIL decision {code} {sep}: Excel={decision} | National={nat_dec} | Word={word_dec}')
        observed[decision] += 1

    word_metrics = {r.cells[0].text: r.cells[1].text for r in doc.tables[2].rows}
    for label, key in [
        ('Total Kasus Di-Review', 'total'), ('Rekomendasi Lanjut On-Site Audit', 'onsite'),
        ('Rekomendasi Lanjut Audit Sampling', 'sampling'), ('Rekomendasi Monitoring / Lolos', 'monitoring'),
        ('Jumlah Perbedaan Dual Coding', 'dc_total'), ('Kasus Mismatch Dual Coding (INA-CBG vs iDRG)', 'dc_cases'),
    ]:
        wval = word_metrics.get(label); sval = h['summary'][key]
        if wval is None or float(wval) != sval:
            errors.append(f'FAIL word summary {code} {key}: {wval!r} != {sval}')

    if snap['snapshot_id'] not in (doc.core_properties.comments or ''):
        errors.append(f'FAIL snapshot_id missing in Word {code}')
    with zipfile.ZipFile(str(d)) as z:
        if any(n.startswith('word/charts/chart') for n in z.namelist()):
            errors.append(f'FAIL old chart data in {code}')

# Word Nasional
doc_nat = Document(str(OUT / 'Laporan_Akhir_Nasional.docx'))
if len(doc_nat.tables[1].rows) - 1 != len(snap['hospitals']):
    errors.append(f'FAIL national Word RS count: {len(doc_nat.tables[1].rows)-1} != {len(snap["hospitals"])}')

# Kontrol Rekonsiliasi Excel Nasional
ctrl = sheet_rows(OUT / 'Rekap_Nasional.xlsx', 'Kontrol Rekonsiliasi')
ctrl_values = [r[1] for r in ctrl[1:11]]
expected_ctrl = [snap['summary'][k] for k in METRICS] + [0]
if ctrl_values != expected_ctrl:
    errors.append(f'FAIL kontrol rekonsiliasi: {ctrl_values} != {expected_ctrl}')

# Tidak ada formula error
for xf in sorted(OUT.rglob('*.xlsx')):
    wb = openpyxl.load_workbook(str(xf), read_only=True, data_only=True)
    for ws in wb:
        for row in ws:
            if any(c.data_type == 'e' for c in row):
                errors.append(f'FAIL Excel formula error in {xf.name} sheet={ws.title}')
    wb.close()

status = 'PASS' if not errors else 'FAIL'
result = dict(
    status=status,
    hospital_count=len(snap['hospitals']),
    case_count=sum(observed.values()),
    decisions=dict(observed),
    summary=snap['summary'],
    snapshot_id=snap['snapshot_id'],
    files={
        'word': len(list(OUT.rglob('*.docx'))),
        'excel': len(list(OUT.rglob('*.xlsx'))),
        'pdf_kkr': len(list((OUT / 'pdf_kkr_per_rs').glob('*.pdf'))) if (OUT / 'pdf_kkr_per_rs').exists() else 0,
    },
    errors=errors[:20],
)
(OUT / 'verifikasi.json').write_text(json.dumps(result, indent=2, ensure_ascii=False), encoding='utf-8')
print(json.dumps(result, indent=2, ensure_ascii=False))
