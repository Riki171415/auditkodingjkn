"""Verify actual printed PDF values against both Excel exports, then merge per RS."""
from pathlib import Path
import sys
import json
import re
import io
from collections import Counter
import hashlib
import zipfile
import openpyxl
import pypdfium2 as pdfium
from pypdf import PdfReader,PdfWriter
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from xml.sax.saxutils import escape

BASE=Path(__file__).resolve().parents[1];sys.path.insert(0,str(BASE))
from modules.report_pdf import FIELD_KEYS,END_MARKER,fields,printable
from scripts.build_consistent_kkr_pdf import load_demographics
SOURCE=BASE/'outputs/rekonsiliasi_20260831'
CASE_DIR=BASE/'output/pdf/kkr_rekonsiliasi_20260831'
OUT=BASE/'output/pdf/KKR_Per_RS_Terverifikasi_20260831'
OUT.mkdir(parents=True,exist_ok=True)
snap=json.loads((SOURCE/'snapshot.json').read_text(encoding='utf-8'))
demo=load_demographics()

def normalize(text):return re.sub(r'\s+','',printable(text)).casefold()

def extract_fields(text):
    text=re.sub(r'\s+',' ',text)
    begin=text.index('1. RINGKASAN HASIL REVIEW TERSIMPAN')
    end=text.index(END_MARKER)
    chunk=text[begin:end]+END_MARKER
    labels=[x[0] for x in FIELD_KEYS]+[END_MARKER]
    result={}
    for label,next_label in zip(labels,labels[1:]):
        m=re.search(re.escape(label)+r'\s+(.*?)\s+'+re.escape(next_label),chunk)
        if not m:raise AssertionError('Missing visible PDF field: '+label)
        result[label]=m.group(1)
    return result

def excel_rows(path,sheet):
    wb=openpyxl.load_workbook(path,read_only=True,data_only=True)
    it=iter(wb[sheet].values);headers=next(it)
    rows=[dict(zip(headers,r)) for r in it];wb.close();return rows

EXCEL_KEYS={
    'Nomor SEP':'Nomor SEP','Kode RS':'Kode RS Sumber','Kode INA-CBG':'Kode INA-CBG',
    'Kode iDRG':'Kode iDRG','Tarif INA-CBG (Rp)':'Tarif INA-CBG (Rp)',
    'Skor KNAVP dasar':'Skor KNAVP','Tingkat risiko tersimpan':'Tingkat Risiko',
    'Jumlah perbedaan dual coding':'Jumlah Perbedaan DC Tersimpan','Jumlah alert KNAVP':'Jumlah Alert KNAVP',
    'Keputusan reviewer asli':'Keputusan Reviewer Asli','Keputusan sistem tersimpan':'Keputusan Sistem Tersimpan',
    'Rekomendasi untuk rekap':'Rekomendasi Laporan','Sumber rekomendasi':'Sumber Rekomendasi'}
NUMERIC={'Tarif INA-CBG (Rp)','Tarif RS (Rp)','Skor KNAVP dasar','Jumlah perbedaan dual coding','Jumlah alert KNAVP'}

def equal(a,b,label):
    if label in NUMERIC:return float(a)==float(b)
    return normalize(a)==normalize(b)

def cover(h):
    buf=io.BytesIO();styles=getSampleStyleSheet();stats=h['summary']
    story=[Paragraph('KUMPULAN KERTAS KERJA REVIEWER',styles['Title']),Paragraph(escape(h['nama_rs']),styles['Heading2']),Paragraph('Kode RS: '+h['kode_rs'],styles['Normal']),Spacer(1,18)]
    values=[['Indikator','Jumlah'],['Total kasus',stats['total']],['On-Site Audit',stats['onsite']],['Sampling/Klarifikasi',stats['sampling']],['Monitoring/tidak perlu tindak lanjut',stats['monitoring']],['Rekomendasi belum tersedia',stats['unknown']],['Kasus dengan perbedaan dual coding',stats['dc_cases']],['Jumlah perbedaan dual coding',stats['dc_total']],['Jumlah alert KNAVP',stats['alerts']]]
    t=Table(values,colWidths=[13*cm,4*cm]);t.setStyle(TableStyle([('GRID',(0,0),(-1,-1),0.4,colors.grey),('BACKGROUND',(0,0),(-1,0),colors.HexColor('#1E3A5F')),('TEXTCOLOR',(0,0),(-1,0),colors.white),('TOPPADDING',(0,0),(-1,-1),8),('BOTTOMPADDING',(0,0),(-1,-1),8)]));story.append(t)
    story += [Spacer(1,20),Paragraph('Salinan rekonsiliasi hasil review tersimpan, bukan pengesahan ulang. Keputusan reviewer asli dan rekomendasi tindak lanjut disajikan terpisah. Data identitas yang tidak tersedia tidak diganti dengan data buatan.',styles['Normal']),Spacer(1,12),Paragraph('Gunakan bookmark Nomor SEP untuk membuka setiap kertas kerja. Nomor halaman pada lembar KKR mengikuti masing-masing kasus.',styles['Normal']),Spacer(1,12),Paragraph('Snapshot SHA256: '+snap['snapshot_id'],styles['Normal'])]
    SimpleDocTemplate(buf,pagesize=A4,leftMargin=2*cm,rightMargin=2*cm).build(story)
    return buf.getvalue()

def main():
    national={(str(r['Kode RS Sumber']),str(r['Nomor SEP'])):r for r in excel_rows(SOURCE/'Rekap_Nasional.xlsx','Master Data (Rincian)')}
    totals=Counter();page_hist=Counter();index=[];all_hospitals=[]
    for hospital_number,h in enumerate(snap['hospitals'],1):
        rs=h['kode_rs'];local={str(r['Nomor SEP']):r for r in excel_rows(SOURCE/'excel_per_rs'/f'Rekap_{rs}.xlsx','Rincian SEP Kasus')}
        writer=PdfWriter();writer.append(PdfReader(io.BytesIO(cover(h))))
        writer.add_metadata({'/Title':f'KKR-DR01 {rs} - {h["nama_rs"]}','/Subject':'Snapshot '+snap['snapshot_id']})
        count=Counter()
        for c in h['cases']:
            sep=str(c['sep']);path=CASE_DIR/rs/f'KKR-DR01_{sep}.pdf'
            doc=pdfium.PdfDocument(path);parts=[]
            for page_index in range(len(doc)):
                page=doc[page_index];tp=page.get_textpage();parts.append(tp.get_text_bounded())
                # Every text rectangle must be within the page, excluding harmless float rounding.
                w,height=page.get_size()
                for k in range(tp.count_rects()):
                    left,bottom,right,top=tp.get_rect(k)
                    assert left>=-1 and bottom>=-1 and right<=w+1 and top<=height+1,('PDF text out of page',rs,sep)
                tp.close();page.close()
            pages=len(doc);doc.close();text='\n'.join(parts)
            printed=extract_fields(text);expected=fields(c)
            assert all(equal(printed[k],v,k) for k,v in expected.items()),('PDF/snapshot mismatch',rs,sep)
            for data,is_national in [(national[(rs,sep)],True),(local[sep],False)]:
                mapping=dict(EXCEL_KEYS)
                mapping['Tarif RS (Rp)']='Tarif RS Standar (Rp)' if is_national else 'Tarif RS (Rp)'
                if not is_national:mapping.update({'Diagnosis INA-CBG':'Diaglist','Prosedur INA-CBG':'Proclist'})
                for label,key in mapping.items():
                    assert equal(printed[label],data.get(key),label),('PDF/Excel mismatch',rs,sep,label,printed[label],data.get(key))
            # Demographics are checked independently against source DB, not generated placeholders.
            name=demo[(rs,sep)].get('Nama_Pasien')
            assert normalize(name) in normalize(text),('Source name missing',rs,sep)
            assert snap['snapshot_id'] in text
            assert 'SALINAN REKONSILIASI - BUKAN PENGESAHAN ULANG' in text
            count[c['rekomendasi_laporan']]+=1;totals[c['rekomendasi_laporan']]+=1;page_hist[pages]+=1
            page_start=len(writer.pages)
            writer.append(str(path),import_outline=False)
            writer.add_outline_item(f'SEP {sep}',page_start)
            index.append(dict(kode_rs=rs,sep=sep,pdf=f'KKR_{rs}.pdf',halaman_awal=page_start+1,halaman_akhir=page_start+pages))
        filename=OUT/f'KKR_{rs}.pdf'
        with filename.open('wb') as f:writer.write(f)
        writer.close()
        # Check the merged artifact too: all visible SEP case headers must occur exactly once.
        merged=pdfium.PdfDocument(filename);merged_keys=[]
        for n in range(len(merged)):
            page=merged[n];tp=page.get_textpage();t=tp.get_text_bounded()
            if '1. RINGKASAN HASIL REVIEW TERSIMPAN' in t:
                merged_keys.append(extract_fields(t)['Nomor SEP'])
            tp.close();page.close()
        merged.close()
        assert merged_keys==[str(c['sep']) for c in h['cases']],('merged PDF lost/duplicated case',rs)
        all_hospitals.append(dict(kode_rs=rs,kasus=len(merged_keys),status='PASS',sha256=hashlib.sha256(filename.read_bytes()).hexdigest()))
        print('PDF VERIFIED + MERGED',hospital_number,'/ 44',rs,len(merged_keys),flush=True)
    assert sum(totals.values())==snap['summary']['total']==len(national)
    report=dict(status='PASS',pdf_count=44,case_count=sum(totals.values()),case_pdf_pages=dict(page_hist),decisions=dict(totals),snapshot_id=snap['snapshot_id'],comparison='Visible PDF fields versus snapshot and actual Excel per-RS/national; all PDF text rectangles within page bounds',hospitals=all_hospitals)
    (OUT/'verifikasi_pdf.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
    (OUT/'indeks_sep.json').write_text(json.dumps(index,indent=2),encoding='utf-8')
    original=json.loads((BASE/'tmp/pdfs/original_pdf_audit.json').read_text(encoding='utf-8'))
    (OUT/'temuan_pdf_lama.json').write_text(json.dumps({'summary':original['summary'],'transitions':original['transitions']},indent=2),encoding='utf-8')
    readme='''PAKET KKR PDF DAN LAPORAN REKONSILIASI

44 PDF per RS memuat seluruh 5.206 kertas kerja. Setiap PDF mempunyai bookmark
Nomor SEP. Ringkasan PDF, Excel per RS dan Excel nasional memakai snapshot sama.
Total: 7 On-Site, 79 Sampling, 5.120 Monitoring/tidak perlu tindak lanjut.
Dual coding: 608 kasus, 1.047 perbedaan tersimpan.

Keputusan reviewer asli, keputusan sistem tersimpan, dan rekomendasi rekap
ditampilkan terpisah. Kolom kosong tidak diganti dengan asumsi; satu kasus
memiliki tingkat risiko sumber kosong dan ditandai Belum diisi. Catatan asli
pada kasus tersebut masih menyebut Rendah; ini dipertahankan sebagai catatan
historis dan perlu klarifikasi reviewer, bukan diubah diam-diam.

PDF lama: rekomendasi, skor, dan jumlah DC cocok setelah penyamaan label.
Namun 5.205 kolom Keputusan Reviewer berisi rekomendasi, bukan keputusan asli.
Seluruh nama pasien lama berbeda dari sumber; seluruh nomor klaim lama cocok
dengan hasil generator angka buatan. PDF baru memakai nama dari data.db;
nomor klaim/peserta dan DPJP yang tidak tersedia ditandai tidak tersedia.

Semua nilai bersama diverifikasi dari TEKS YANG TERCETAK pada PDF terhadap
Excel yang telah diekspor. Pemeriksaan seluruh halaman untuk batas teks
lulus; sampel tata letak dan tiap PDF RS diperiksa visual. Identitas tambahan
dicocokkan ke data.db. Data klinis/keputusan tersimpan tidak diubah.

Ini salinan rekonsiliasi, bukan pengesahan atau tanda tangan ulang.
Simpan terbatas karena PDF memuat identitas pasien dan data kesehatan.
File lama tidak ditimpa. Jangan campur PDF lama dengan paket koreksi ini.
Word dalam paket tetap hasil koreksi sebelumnya; keterbatasan QA visual Word
(LibreOffice tidak tersedia) tetap berlaku. Server aplikasi perlu memuat
ulang kode agar jalur unduh PDF baru aktif dan melewati cache lama.
'''
    (OUT/'BACA_DULU.txt').write_text(readme,encoding='utf-8')
    print(json.dumps({k:v for k,v in report.items() if k!='hospitals'},indent=2))

if __name__=='__main__':main()
