"""Read finished Word/Excel files independently; fail on any reconciliation error."""
import json
from collections import Counter
from pathlib import Path
import sys
import zipfile
from docx import Document
import openpyxl

BASE=Path(__file__).resolve().parents[1]
OUT=BASE/'outputs/rekonsiliasi_20260831'
snap=json.loads((OUT/'snapshot.json').read_text(encoding='utf-8'))

def rows(path,sheet):
    wb=openpyxl.load_workbook(path,read_only=True,data_only=True)
    data=list(wb[sheet].values)
    wb.close()
    return data

metrics=['total','onsite','sampling','monitoring','unknown','dc_cases','dc_total','alerts','avg_score']
national=rows(OUT/'Rekap_Nasional.xlsx','Ringkasan Eksekutif')
national_by_rs={str(r[0]):r for r in national[1:]}
national_detail=rows(OUT/'Rekap_Nasional.xlsx','Master Data (Rincian)')
hdr=national_detail[0]
national_cases={(str(r[hdr.index('Kode RS Sumber')]),str(r[hdr.index('Nomor SEP')])):r[hdr.index('Rekomendasi Laporan')] for r in national_detail[1:]}
observed=Counter()
for h in snap['hospitals']:
    code=h['kode_rs']
    x=OUT/'excel_per_rs'/f'Rekap_{code}.xlsx'
    d=OUT/'word_per_rs'/f'LHR_{code}.docx'
    summary=rows(x,'Ringkasan Eksekutif RS')[1]
    expected=[h['summary'][k] for k in metrics]
    assert list(summary[2:11])==expected,(code,summary,expected)
    assert list(national_by_rs[code][2:11])==expected,('national row',code)
    assert summary[11]==0
    detail=rows(x,'Rincian SEP Kasus'); heads=detail[0]
    doc=Document(d)
    narrative='\n'.join(p.text for p in doc.paragraphs)
    expected_narrative=(f'Dari {h["summary"]["total"]} kasus: '
        f'{h["summary"]["onsite"]} direkomendasikan On-Site Audit, '
        f'{h["summary"]["sampling"]} Sampling/Klarifikasi, '
        f'{h["summary"]["monitoring"]} Monitoring/tidak perlu tindak lanjut, dan '
        f'{h["summary"]["unknown"]} belum memiliki rekomendasi yang dapat diklasifikasikan.')
    assert expected_narrative in narrative, ('word narrative',code)
    assert 'Hasil penilaian risiko menunjukkan bahwa' not in narrative, ('stale narrative',code)
    assert len(doc.tables[1].rows)-1==h['summary']['total'],('word count',code)
    word_cases={r.cells[1].text:r.cells[12].text for r in doc.tables[1].rows[1:]}
    for r in detail[1:]:
        sep=str(r[heads.index('Nomor SEP')]);decision=r[heads.index('Rekomendasi Laporan')]
        assert decision==national_cases[(code,sep)]==word_cases[sep],('decision',code)
        observed[decision]+=1
    word_metrics={r.cells[0].text:r.cells[1].text for r in doc.tables[2].rows}
    for label,key in [('Total Kasus Di-Review','total'),('Rekomendasi Lanjut On-Site Audit','onsite'),('Rekomendasi Lanjut Audit Sampling','sampling'),('Rekomendasi Monitoring / Lolos','monitoring'),('Jumlah Perbedaan Dual Coding','dc_total'),('Kasus Mismatch Dual Coding (INA-CBG vs iDRG)','dc_cases')]:
        assert float(word_metrics[label])==h['summary'][key],('word summary',code,key)
    expected_priority=sum(1 for c in h['cases'] if c['jumlah_beda_dual_coding'] > 0 or len(c['triggered_rules']) > 0)
    assert len(doc.tables[4].rows)-1==expected_priority, ('word priority',code)
    assert snap['snapshot_id'] in doc.core_properties.comments
    with zipfile.ZipFile(d) as z:
        assert not any(n.startswith('word/charts/chart') for n in z.namelist()),'Old chart data remains'
doc=Document(OUT/'Laporan_Akhir_Nasional.docx')
national_narrative='\n'.join(p.text for p in doc.paragraphs)
s=snap['summary']
fmt=lambda value:f'{value:,}'.replace(',','.')
assert (f'Dari {fmt(s["total"])} kasus: {fmt(s["onsite"])} direkomendasikan On-Site Audit, '
        f'{fmt(s["sampling"])} Sampling/Klarifikasi, {fmt(s["monitoring"])} Monitoring/tidak perlu tindak lanjut, dan '
        f'{fmt(s["unknown"])} belum memiliki rekomendasi yang dapat diklasifikasikan.') in national_narrative
assert 'Hasil penilaian risiko menunjukkan bahwa' not in national_narrative
assert len(doc.tables[1].rows)-1==len(snap['hospitals'])
for r in doc.tables[1].rows[1:]:
    c=[x.text for x in r.cells];s=next(h['summary'] for h in snap['hospitals'] if h['kode_rs']==c[1])
    assert [float(v) for v in c[3:]]==[s[k] for k in ('total','onsite','sampling','monitoring','unknown','dc_cases','dc_total','alerts','avg_score','cases_with_alerts')]
expected=[snap['summary'][k] for k in metrics]+[0]
assert [r[1] for r in rows(OUT/'Rekap_Nasional.xlsx','Kontrol Rekonsiliasi')[1:11]]==expected
for file in OUT.rglob('*.xlsx'):
    wb=openpyxl.load_workbook(file,read_only=True,data_only=True)
    for ws in wb:
        for row in ws:
            assert not any(c.data_type=='e' for c in row),(file.name,ws.title,'Excel error')
    wb.close()
assert len(list(OUT.rglob('*.docx')))==45, 'Unexpected Word duplicate or missing file'
result={'status':'PASS','hospital_count':len(snap['hospitals']),'case_count':sum(observed.values()),'decisions':dict(observed),'summary':snap['summary'],'snapshot_id':snap['snapshot_id'],'files':{'word':len(list(OUT.rglob('*.docx'))),'excel':len(list(OUT.rglob('*.xlsx')))},'narrative_check':'PASS: every per-RS and national narrative is generated from the same snapshot','visual_word_qa':'Not completed: LibreOffice unavailable'}
(OUT/'verifikasi.json').write_text(json.dumps(result,indent=2,ensure_ascii=False),encoding='utf-8')
print(json.dumps(result,indent=2))
