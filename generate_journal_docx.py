import os
import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def add_heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    for run in h.runs:
        run.font.name = 'Times New Roman'
        run.font.color.rgb = RGBColor(0, 0, 0)
        run.font.bold = True
        run.font.size = Pt(14) if level == 1 else Pt(12)
    return h

def main():
    md_path = 'exports/Final_Manuscript_Jurnal.md'
    out_path = 'exports/Final_Manuscript_Jurnal.docx'
    
    if not os.path.exists(md_path):
        print("MD file not found!")
        return
        
    with open(md_path, 'r', encoding='utf-8') as f:
        md_text = f.read()

    doc = Document()
    
    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    
    # Set double spacing for academic standard
    style.paragraph_format.line_spacing = 2.0
    
    lines = md_text.split('\n')
    for line in lines:
        line = line.strip()
        if not line or line == '---' or line == '*End of Manuscript Draft*':
            continue
            
        if line.startswith('# '):
            add_heading(doc, line[2:].strip(), 1)
        elif line.startswith('## '):
            add_heading(doc, line[3:].strip(), 2)
        elif line.startswith('### '):
            add_heading(doc, line[4:].strip(), 3)
        elif line.startswith('- ') or line.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            text = line[2:].strip()
            parts = re.split(r'(\*\*.*?\*\*|\*.*?\*)', text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    p.add_run(part[2:-2]).bold = True
                elif part.startswith('*') and part.endswith('*'):
                    p.add_run(part[1:-1]).italic = True
                else:
                    p.add_run(part)
        elif re.match(r'^\d+\.\s', line):
            p = doc.add_paragraph(style='List Number')
            text = re.sub(r'^\d+\.\s+', '', line)
            parts = re.split(r'(\*\*.*?\*\*|\*.*?\*)', text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    p.add_run(part[2:-2]).bold = True
                elif part.startswith('*') and part.endswith('*'):
                    p.add_run(part[1:-1]).italic = True
                else:
                    p.add_run(part)
        else:
            p = doc.add_paragraph()
            parts = re.split(r'(\*\*.*?\*\*|\*.*?\*)', line)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    p.add_run(part[2:-2]).bold = True
                elif part.startswith('*') and part.endswith('*'):
                    p.add_run(part[1:-1]).italic = True
                elif part.startswith('<sup>') and part.endswith('</sup>'):
                    run = p.add_run(part[5:-6])
                    run.font.superscript = True
                else:
                    # simplistic check for superscripts mixed in text like [Author Name]<sup>1</sup>
                    subparts = re.split(r'(<sup>.*?</sup>)', part)
                    for sub in subparts:
                        if sub.startswith('<sup>') and sub.endswith('</sup>'):
                            run = p.add_run(sub[5:-6])
                            run.font.superscript = True
                        else:
                            p.add_run(sub)

    doc.save(out_path)
    print(f"DOCX created at {out_path}")

if __name__ == '__main__':
    main()
