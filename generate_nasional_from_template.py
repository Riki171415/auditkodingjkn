import os
import re
import json
import shutil
import io
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_ORIENT
from modules.db_manager import get_recap_desk_review

from modules.export_generator import generate_qr_image_bytes as generate_custom_qr_bytes, QR_AVAILABLE, PIL_AVAILABLE

def set_normal_style(run, size=12):
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)

def _add_formatted_runs(p, text):
    parts = re.split(r'(\*\*.*?\*\*|\*.*?\*|\[.*?\]\(.*?\))', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = p.add_run(part[2:-2])
            run.bold = True
            set_normal_style(run)
        elif part.startswith('*') and part.endswith('*'):
            run = p.add_run(part[1:-1])
            run.italic = True
            set_normal_style(run)
        elif part.startswith('[') and '](' in part and part.endswith(')'):
            clean_text = re.sub(r'\[(.+?)\]\(.+?\)', r'\1', part)
            run = p.add_run(clean_text)
            set_normal_style(run)
        else:
            run = p.add_run(part)
            set_normal_style(run)

def parse_markdown_blocks(md_text):
    blocks = []
    lines = md_text.split('\n')
    i = 0
    in_toc = False
    while i < len(lines):
        line = lines[i]
        if '<!-- TOC_START -->' in line:
            in_toc = True; i += 1; continue
        if '<!-- TOC_END -->' in line:
            in_toc = False; i += 1; continue
        if in_toc:
            i += 1; continue
            
        if line.startswith('#'):
            level = len(line) - len(line.lstrip('#'))
            header_text = line.lstrip('#').strip()
            blocks.append({'type': 'header', 'level': level, 'content': header_text})
        elif line.strip() and not re.match(r'^[-*_]{3,}$', line.strip()):
            blocks.append({'type': 'text', 'content': line.strip()})
        i += 1
    return blocks

def generate_nasional_report():
    from modules.report_data import load_snapshot
    from modules.report_word import write_report
    snapshot = load_snapshot()
    out_path = os.path.join('exports', 'word_reports', 'Laporan_Akhir_Nasional_V2.docx')
    return write_report(snapshot['cases'], out_path, snapshot['snapshot_id'], snapshot['hospitals'])


def _generate_nasional_report_legacy():
    print("Membuat Laporan Nasional Berbasis Template dengan 6 Lampiran...")
    
    md_path = os.path.join('exports', 'agent_outputs', 'final_md', 'final_md_nasional.md')
    if not os.path.exists(md_path):
        print(f"File markdown {md_path} tidak ditemukan!")
        return
    with open(md_path, 'r', encoding='utf-8') as f:
        md_text = f.read()
        
    blocks = parse_markdown_blocks(md_text)
    
    OUTPUT_DIR = os.path.join('exports', 'word_reports')
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    out_path = os.path.join(OUTPUT_DIR, "Laporan_Akhir_Nasional_V2.docx")
    
    shutil.copy('Template_Laporan_RS.docx', out_path)
    doc = Document(out_path)
    
    def iter_all_paragraphs(document):
        for p in document.paragraphs: yield p
        for t in document.tables:
            for r in t.rows:
                for c in r.cells:
                    for p in c.paragraphs: yield p

    for p in iter_all_paragraphs(doc):
        if '{NAMA_RS}' in p.text:
            p.text = p.text.replace('{NAMA_RS}', 'Nasional (Agregat 44 RS)')
        if '{KODE_RS}' in p.text:
            p.text = p.text.replace('{KODE_RS}', 'NASIONAL-001')
        if '1275655' in p.text:
            p.text = p.text.replace('1275655', 'Nasional')
        if 'RSU H ADAM MALIK' in p.text:
            p.text = p.text.replace('RSU H ADAM MALIK', 'Nasional')

    for p in doc.paragraphs:
        if '{KONTEN_LAPORAN}' in p.text:
            p.text = '' 
            for block in blocks:
                if block['type'] == 'header':
                    new_p = p.insert_paragraph_before('')
                    if block['content'].startswith('BAB '):
                        new_p.paragraph_format.page_break_before = True
                    run = new_p.add_run(block['content'])
                    run.bold = True
                    run.font.name = 'Times New Roman'
                    if block['level'] == 1:
                        run.font.size = Pt(14)
                        new_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    else:
                        run.font.size = Pt(12)
                elif block['type'] == 'text':
                    b_content = block['content']
                    if '[CHART_' in b_content:
                        continue # Skip placeholder, we use native charts
                    if b_content.startswith('- ') or b_content.startswith('* '):
                        new_p = p.insert_paragraph_before('')
                        _add_formatted_runs(new_p, b_content)
                        new_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                        new_p.paragraph_format.left_indent = Pt(36)
                        new_p.paragraph_format.first_line_indent = Pt(-18)
                    elif re.match(r'^\d+\.\s', b_content):
                        new_p = p.insert_paragraph_before('')
                        _add_formatted_runs(new_p, b_content)
                        new_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                        new_p.paragraph_format.left_indent = Pt(36)
                        new_p.paragraph_format.first_line_indent = Pt(-18)
                    else:
                        new_p = p.insert_paragraph_before('')
                        _add_formatted_runs(new_p, b_content)
                        new_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                        new_p.paragraph_format.first_line_indent = Pt(36)
                        
    # Populate Tables
    recap = get_recap_desk_review()
    
    def _risk_w(c):
        tr = str(c.get('tingkat_risiko', '')).lower()
        if tr == 'tinggi': return 3
        if tr == 'sedang': return 2
        return 1
        
    recap_sorted = sorted(recap, key=lambda x: (-_risk_w(x), x.get('sep', '')))
    
    # Table 1: Ringkasan KKR-DR01 (Top 200)
    t1 = doc.tables[1]
    top_200 = recap_sorted[:200]
    for idx, case in enumerate(top_200, 1):
        try: fd = json.loads(case.get('tindakan_reviewer') or '{}')
        except: fd = {}
        row = t1.add_row().cells
        row[0].text = str(idx)
        row[1].text = str(case.get('sep', ''))
        row[2].text = str(case.get('diaglist', ''))
        row[3].text = str(case.get('proclist', ''))
        row[4].text = str(case.get('inacbg', ''))
        row[5].text = str(case.get('deskripsi_inacbg', ''))
        row[6].text = str(case.get('idrg_code', ''))
        row[7].text = str(case.get('deskripsi_idrg', ''))
        
        rules = case.get('triggered_rules', [])
        skor = sum(int(r.get('bobot', 1)) for r in rules)
        row[8].text = str(skor)
        row[9].text = str(case.get('tingkat_risiko', ''))
        
        if rules:
            rule_parts = [r.get('nama_aturan', '') for r in rules if r.get('nama_aturan')]
            keterangan = ", ".join(rule_parts)
        else:
            keterangan = "-"
        row[10].text = keterangan
        beda_dc = int(fd.get('jumlah_beda_dual_coding', 0) or 0)
        row[11].text = "Ada Beda" if beda_dc > 0 else "Sesuai"
        row[12].text = str(fd.get('keputusan_sistem', '-'))
        
        for c in row:
            for r in c.paragraphs[0].runs: set_normal_style(r, 7)
            
    # Table 2: Rekapitulasi Hasil Validasi KNAVP
    t2 = doc.tables[2]
    kep_counts = {'Direkomendasikan On-Site Audit': 0, 'Audit Sampling': 0, 'Lolos/Monitoring': 0}
    severity_counts = {'High': 0, 'Medium': 0, 'Low': 0}
    discrepancy_count = 0
    total_skor = 0
    
    for r in recap:
        try: fd = json.loads(r.get('tindakan_reviewer') or '{}')
        except: fd = {}
        kep = fd.get('keputusan_sistem', 'Lolos/Monitoring')
        if kep in kep_counts: kep_counts[kep] += 1
        else: kep_counts['Lolos/Monitoring'] += 1
        
        if int(fd.get('jumlah_beda_dual_coding', 0) or 0) > 0: discrepancy_count += 1
        
        rules = r.get('triggered_rules', [])
        skor = sum(int(ru.get('bobot', 1)) for ru in rules)
        total_skor += skor
        
        if rules:
            for ru in rules:
                sev = str(ru.get('severity', 'low')).lower()
                if sev == 'high': severity_counts['High'] += 1
                elif sev == 'medium': severity_counts['Medium'] += 1
                else: severity_counts['Low'] += 1
                
    total_cases = len(recap)
    t2.rows[0].cells[1].text = str(total_cases)
    t2.rows[1].cells[1].text = str(sum(severity_counts.values()))
    t2.rows[2].cells[1].text = str(severity_counts['High'])
    t2.rows[3].cells[1].text = str(severity_counts['Medium'])
    t2.rows[4].cells[1].text = str(severity_counts['Low'])
    t2.rows[5].cells[1].text = str(round(total_skor / total_cases, 2)) if total_cases > 0 else "0"
    t2.rows[6].cells[1].text = str(kep_counts['Direkomendasikan On-Site Audit'])
    t2.rows[7].cells[1].text = str(kep_counts['Audit Sampling'])
    t2.rows[8].cells[1].text = str(kep_counts['Lolos/Monitoring'])
    t2.rows[9].cells[1].text = str(discrepancy_count)
    t2.rows[10].cells[1].text = f"{round((discrepancy_count/total_cases*100), 1)}%" if total_cases > 0 else "0%"
    
    # Table 3: Detail Hasil Audit KNAVP per Aturan
    rule_map = {}
    for r in recap:
        rules = r.get('triggered_rules', [])
        for ru in rules:
            rid = ru.get('rule_id', '-')
            if rid not in rule_map:
                rule_map[rid] = {
                    'Kategori': ru.get('kelompok_rule', '-'),
                    'Nama Aturan': ru.get('nama_aturan', '-'),
                    'Severity': ru.get('severity', 'Low'),
                    'count': 0
                }
            rule_map[rid]['count'] += 1
            
    t3 = doc.tables[3]
    for idx, (rid, data) in enumerate(sorted(rule_map.items(), key=lambda x: x[1]['count'], reverse=True), 1):
        row = t3.add_row().cells
        row[0].text = str(idx)
        row[1].text = str(data['Kategori'])
        row[2].text = str(rid)
        row[3].text = str(data['Nama Aturan'])
        row[4].text = str(data['Severity'])
        row[5].text = str(data['count']) + " kasus"
        row[6].text = "-"
        row[7].text = "-"
        row[8].text = "-"
        for c in row:
            for r in c.paragraphs[0].runs: set_normal_style(r, 9)

    # Table 4: Daftar Kasus Prioritas
    t4 = doc.tables[4]
    onsite_cases = [c for c in recap_sorted if 'On-Site' in str(c.get('keputusan', '')) or 'On-Site' in str(json.loads(c.get('tindakan_reviewer') or '{}').get('keputusan_sistem', ''))]
    for idx, case in enumerate(onsite_cases[:100], 1):
        try: fd = json.loads(case.get('tindakan_reviewer') or '{}')
        except: fd = {}
        row = t4.add_row().cells
        row[0].text = str(idx)
        row[1].text = str(case.get('sep', ''))
        row[2].text = str(case.get('inacbg', ''))
        row[3].text = str(case.get('deskripsi_inacbg', ''))
        row[4].text = str(case.get('diaglist', ''))
        row[5].text = str(case.get('proclist', ''))
        rules = case.get('triggered_rules', [])
        skor = sum(int(r.get('bobot', 1)) for r in rules)
        row[6].text = str(skor)
        row[7].text = str(case.get('tingkat_risiko', ''))
        row[8].text = str(len(rules))
        beda_dc = int(fd.get('jumlah_beda_dual_coding', 0) or 0)
        row[9].text = "Ada Beda" if beda_dc > 0 else "Sesuai"
        
        if rules:
            rule_parts = [r.get('nama_aturan', '') for r in rules if r.get('nama_aturan')]
            keterangan = ", ".join(rule_parts)
        else:
            keterangan = "-"
        row[10].text = keterangan
        row[11].text = str(fd.get('keputusan_sistem', '-'))
        
        for c in row:
            for r in c.paragraphs[0].runs: set_normal_style(r, 8)

    # Signature Block
    doc.add_page_break()
    t_sig = doc.add_table(rows=2, cols=3)
    t_sig_hdr = t_sig.rows[0].cells
    t_sig_hdr[0].text = 'Disusun oleh'
    t_sig_hdr[1].text = 'Direviu oleh'
    t_sig_hdr[2].text = 'Keterangan'
    
    penyusun_name = "Tim Reviewer Koding PUSBIKES"
    ketua_name = "Riki Permana Putra.,SKM"
    ketua_nip = "198611172014021001"

    t_sig_val = t_sig.rows[1].cells
    t_sig_val[0].text = ''
    p_rev = t_sig_val[0].paragraphs[0]
    p_rev.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if QR_AVAILABLE and PIL_AVAILABLE:
        try:
            rev_payload = json.dumps({"role": "Penyusun", "name": penyusun_name}, ensure_ascii=False)
            r_bytes = generate_custom_qr_bytes(rev_payload, size_px=130)
            if r_bytes:
                run_r = p_rev.add_run()
                run_r.add_picture(io.BytesIO(r_bytes), width=Inches(1.2))
        except Exception: pass
            
    p_rev_name = t_sig_val[0].add_paragraph(f"( {penyusun_name} )")
    p_rev_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p_rev_name.runs: run.font.name = 'Times New Roman'; run.bold = True; run.font.size = Pt(9.5)

    t_sig_val[1].text = ''
    p_qr = t_sig_val[1].paragraphs[0]
    p_qr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if QR_AVAILABLE and PIL_AVAILABLE:
        try:
            k_payload = json.dumps({"role": "Ketua", "name": ketua_name, "nip": ketua_nip}, ensure_ascii=False)
            k_bytes = generate_custom_qr_bytes(k_payload, size_px=130)
            if k_bytes:
                run_qr = p_qr.add_run()
                run_qr.add_picture(io.BytesIO(k_bytes), width=Inches(1.2))
        except Exception: pass

    p_name = t_sig_val[1].add_paragraph(f"( {ketua_name} )\nNIP. {ketua_nip}")
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p_name.runs: 
        run.font.name = 'Times New Roman'
        run.bold = True if 'NIP' not in run.text else False
        run.font.size = Pt(9.5) if 'NIP' in run.text else Pt(10)

    t_sig_val[2].text = ''
    p_ket = t_sig_val[2].paragraphs[0]
    p_ket.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_ket_run = p_ket.add_run("1. Laporan disusun oleh Tim Reviewer Koding PUSBIKES.\n2. Barcode QR memvalidasi dokumen.\n3. Dokumen ini sah sebagai rekapitulasi Desk Review Nasional 2025.")
    p_ket_run.font.name = 'Times New Roman'
    p_ket_run.font.size = Pt(9.5)
    
    t_sig.style = 'Table Grid'
    
    try:
        doc.save(out_path)
        print(f"BERHASIL: Laporan Akhir Nasional tersimpan di: {out_path}")
    except PermissionError:
        out_path = out_path.replace('.docx', '_Baru.docx')
        doc.save(out_path)
        print(f"BERHASIL: Laporan Akhir Nasional tersimpan di: {out_path}")

if __name__ == '__main__':
    generate_nasional_report()
