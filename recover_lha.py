import re
import sys

def reconstruct():
    # 1. Read the git version of export_generator.py
    with open('modules/export_generator.py', 'r', encoding='utf-8') as f:
        git_content = f.read()
    
    # 2. Find where generate_lha_word starts
    match = re.search(r'def generate_lha_word\(', git_content)
    if not match:
        print("Could not find generate_lha_word")
        return
        
    start_pos = match.start()
    git_base = git_content[:start_pos]
    
    # 3. Read the temp files
    def get_text(filename):
        with open(filename, 'r', encoding='utf-8') as f:
            return f.read()
            
    katpeng2 = get_text('temp_katpeng2.txt')
    # Actually, we need the initialization of document before katpeng2.
    # katpeng2 starts with `["Rumah Sakit", f":  {rs_name}"],`
    init_code = """def generate_lha_word(kode_rs, rs_name, cases, output_path, summary_stats=None):
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    import os

    document = Document()

    def add_p(text, align=WD_ALIGN_PARAGRAPH.LEFT, bold=False, size_pt=10, space_after_pt=0):
        p = document.add_paragraph()
        p.alignment = align
        if space_after_pt:
            p.paragraph_format.space_after = Pt(space_after_pt)
        run = p.add_run(text)
        run.bold = bold
        run.font.name = 'Times New Roman'
        run.font.size = Pt(size_pt)
        return p

    def add_list_item(bullet, text, size_pt=10, space_after_pt=0):
        p = document.add_paragraph()
        p.style = 'List Bullet'
        p.paragraph_format.space_after = Pt(space_after_pt)
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(size_pt)
        return p

    def format_table(table, font_size_pt=10):
        table.style = 'Table Grid'
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = 'Times New Roman'
                        run.font.size = Pt(font_size_pt)
                        
    # ── Title ─────────────────────────────────────────────────────────────────
    heading = document.add_heading('LAPORAN HASIL DESK REVIEW AUDIT CODING', 0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = document.add_paragraph('Transisi INA-CBG menuju Indonesian Diagnosis Related Groups (iDRG)')
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].italic = True
    
    document.add_paragraph("", style='Normal')
    
    table_meta = document.add_table(rows=6, cols=2)
    meta_rows = [
"""
    
    # 4. Now we piece them together.
    # We will build the new function string
    new_func = init_code + katpeng2 
    
    # temp_katpeng.txt starts after Kata Pengantar
    katpeng = get_text('temp_katpeng.txt')
    new_func += katpeng
    
    # temp_bab1.txt has BAB I and BAB II A, B
    bab1 = get_text('temp_bab1.txt')
    new_func += bab1
    
    # temp_analisis.txt has BAB II C
    analisis = get_text('temp_analisis.txt')
    new_func += analisis
    
    # temp_bab2_d.txt has BAB II D
    bab2_d = get_text('temp_bab2_d.txt')
    new_func += bab2_d
    
    # Add the rest of BAB II D (loop for table_kp)
    bab2_d_rest = """
    priority_cases = cases_with_rules + cases_dc_only
    
    for i, c in enumerate(priority_cases):
        row_cells = table_kp.add_row().cells
        row_cells[0].text = str(i + 1)
        row_cells[1].text = str(c.get('sep', ''))
        row_cells[2].text = str(c.get('triggered_rules', 'Perbedaan Dual Coding Tinggi'))
        row_cells[3].text = str(c.get('tingkat_risiko', 'Tinggi'))
        row_cells[4].text = "On-Site Audit / Klarifikasi"
    format_table(table_kp, font_size_pt=10)
    add_p("", space_after_pt=12)
"""
    new_func += bab2_d_rest
    
    # 5. Add BAB II E (Kesesuaian Input)
    bab2_e = """
    document.add_heading('E. Kesesuaian Input INACBG-iDRG (Discrepancy Koding)', level=2)
    
    total_cases = len(cases)
    mismatch_cases = len([c for c in cases if int(c.get('jumlah_beda_dual_coding', 0)) > 0])
    match_cases = total_cases - mismatch_cases
    mismatch_pct = round((mismatch_cases / total_cases * 100), 1) if total_cases else 0
    match_pct = round((match_cases / total_cases * 100), 1) if total_cases else 0
    tot_beda = sum(int(c.get('jumlah_beda_dual_coding', 0) or 0) for c in cases)
    
    narasi_e = f"Berdasarkan hasil evaluasi terhadap {total_cases} tagihan pasien di {rs_name}, ditemukan sebanyak {mismatch_cases} tagihan atau {mismatch_pct}% yang mengalami ketidaksesuaian antara input diagnosis atau tindakan pada sistem INACBG dan iDRG, dengan total keseluruhan mencapai {tot_beda} perbedaan kode medis. Sebaliknya, sebanyak {match_cases} tagihan atau {match_pct}% sisanya telah menunjukkan kesesuaian data input antara kedua sistem tersebut."
    add_p(narasi_e, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size_pt=12, space_after_pt=8)
    
    add_p("Berikut adalah profil ringkas akurasi pengodean medis di rumah sakit ini:", size_pt=12, space_after_pt=4)
    
    bullets_e = [
        f"Kasus yang Sudah Cocok ({match_pct}% atau {match_cases} pasien):\\nTagihan ini aman dan sinkron. Kode ICD yang dikirim ke INA-CBG sudah sama persis dengan aturan iDRG.",
        f"Kasus yang Perlu Dicek Ulang ({mismatch_pct}% atau {mismatch_cases} pasien):\\nPada kelompok ini, sistem iDRG membaca adanya kode yang hilang, berlebih, atau bergeser dibandingkan klaim awal INA-CBG. Total ada {tot_beda} selisih kode (terdiri dari tambahan maupun pengurangan diagnosis/tindakan).",
        f"Potensi Perubahan (0.0%):\\nDari seluruh kasus yang tidak sinkron tersebut, belum ditemukan indikasi pergeseran tingkat keparahan (CCL/PCCL) yang akan berdampak langsung pada perubahan tarif dasar."
    ]
    for b in bullets_e:
        bold_part, rest = b.split(':\\n', 1)
        p = add_list_item("●", "")
        r1 = p.add_run(bold_part + ": ")
        r1.bold = True
        p.add_run(rest)
        
    add_p("(Catatan: Rincian selisih kode per kasus dapat dilihat pada Lampiran 1: KKR-DR01)", size_pt=10, bold=True, space_after_pt=12)
"""
    new_func += bab2_e
    
    # 6. Add BAB III and Lampiran
    lampiran = get_text('temp_lampiran.txt')
    # temp_lampiran.txt starts with some remaining `add_p` from the old BAB II E, we should slice it to start from `document.add_page_break()`
    match = re.search(r'document\.add_page_break\(\)\s*# ── BAB III', lampiran)
    if match:
        lampiran = lampiran[match.start():]
        
    # Replace Lampiran 2 and 4 tables with bullets in the lampiran string
    match_l2 = re.search(r"document\.add_heading\('2\. Rekapitulasi Hasil Validasi KNAVP', level=2\)(.*?)add_p\(\"\", space_after_pt=6\)", lampiran, re.DOTALL)
    if match_l2:
        new_l2 = """document.add_heading('2. Rekapitulasi Hasil Validasi KNAVP', level=2)
    ss = summary_stats or {}
    onsite   = ss.get('onsite',   sum(1 for c in cases if 'On-Site' in str(c.get('keputusan_sistem', ''))))
    sampling = ss.get('sampling', sum(1 for c in cases if 'Sampling' in str(c.get('keputusan_sistem', ''))))
    monitor  = ss.get('monitoring', len(cases) - onsite - sampling)
    avg_skor = ss.get('avg_skor_knavp', round(sum(float(c.get('knavp_skor', 0) or 0) for c in cases) / max(len(cases), 1), 1))
    beda_dc  = ss.get('total_beda_dc', sum(int(c.get('jumlah_beda_dual_coding', 0) or 0) for c in cases))

    l2_bullets = [
        f"Total Kasus Di-Review: {len(cases)}",
        f"Rekomendasi Lanjut On-Site Audit: {onsite}",
        f"Rekomendasi Lanjut Sampling: {sampling}",
        f"Rekomendasi Monitoring (Lolos): {monitor}",
        f"Rata-rata Skor KNAVP: {avg_skor}",
        f"Total Perbedaan Dual Coding (INA-CBG vs iDRG): {beda_dc}",
        f"Status Pengelompokan iDRG: Terverifikasi Rule Engine"
    ]
    for b in l2_bullets:
        add_list_item("●", b)
    add_p("", space_after_pt=6)"""
        lampiran = lampiran[:match_l2.start()] + new_l2 + lampiran[match_l2.end():]

    match_l4 = re.search(r"document\.add_heading\('4\. Dashboard Hasil Validasi.*?level=2\)(.*?)(?=else:\s*add_p\(\"Tidak ada|document\.save)", lampiran, re.DOTALL)
    if match_l4:
        new_l4 = """document.add_heading('4. Dashboard Hasil Validasi (Executive Summary KPI)', level=2)
    compliance_rate = (monitor / len(cases) * 100) if cases else 100
    outlier_rate = (onsite / len(cases) * 100) if cases else 0
    risk_class = "Tinggi" if avg_skor >= 4 else "Sedang" if avg_skor >= 2 else "Rendah"
    
    l4_bullets = [
        f"Tingkat Kepatuhan Koding (Compliance Rate): {round(compliance_rate,1)}% ({monitor} kasus lolos tanpa catatan/risiko minor)",
        f"Rasio Outlier (On-Site Audit Rate): {round(outlier_rate,1)}% ({onsite} kasus wajib diverifikasi fisik rekam medis)",
        f"Indeks Skor Risiko KNAVP RS: {avg_skor} / 10 (Klasifikasi Risiko {risk_class})"
    ]
    for b in l4_bullets:
        add_list_item("●", b)
        
    add_p("Distribusi Status & Rekomendasi Kasus:", bold=True, space_after_pt=4)
    samp_rate = (sampling / len(cases) * 100) if cases else 0
    dist_bullets = [
        f"Monitoring (Lolos / Risiko Rendah): {monitor} Kasus ({round(compliance_rate,1)}%)",
        f"Sampling / Klarifikasi (Risiko Sedang): {sampling} Kasus ({round(samp_rate,1)}%)",
        f"On-Site Audit (Risiko Tinggi ≥ 4.0): {onsite} Kasus ({round(outlier_rate,1)}%)",
        f"TOTAL KASUS SAMPEL: {len(cases)} Kasus (100%)"
    ]
    for b in dist_bullets:
        add_list_item("●", b)
    add_p("", space_after_pt=12)
    
    """
        lampiran = lampiran[:match_l4.start()] + new_l4 + lampiran[match_l4.end():]
        
    # Also we need to make sure `document.save(output_path)` and `return output_path` are at the end
    # They should be there because match_l4 looked until `document.save`. But let's check.
    
    if "document.save(output_path)" not in lampiran:
        lampiran += "\\n    document.save(output_path)\\n    return output_path\\n"
        
    new_func += lampiran
    
    # 7. Write the new export_generator.py
    final_content = git_base + "\\n\\n" + new_func
    with open('modules/export_generator.py', 'w', encoding='utf-8') as f:
        f.write(final_content)
        
    print("Reconstructed export_generator.py successfully!")

if __name__ == '__main__':
    reconstruct()
