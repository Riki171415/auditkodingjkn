import os
import json
import hashlib
import random
import string
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from modules.db_manager import get_recap_desk_review

OUTPUT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'exports', 'rekap_per_rs')

CCL_MAP = {
    '0': 'Tanpa Komplikasi (0)',
    '1': 'Komplikasi Ringan (1)',
    '2': 'Komplikasi Sedang (2)',
    '3': 'Komplikasi Berat (3)',
    '4': 'Komplikasi Sangat Berat (4)',
}

# Palet warna premium
BLUE_DARK  = '1E3A5F'
WHITE      = 'FFFFFF'
LIGHT_BLUE = 'DBEAFE'
LIGHT_RED  = 'FEF2F2'
LIGHT_AMB  = 'FFFBEB'
LIGHT_GREEN = 'F0FFF4'

thin = Side(style='thin')
border = Border(left=thin, right=thin, top=thin, bottom=thin)

def _hdr_style(fill_color=BLUE_DARK):
    return {
        'fill': PatternFill(start_color=fill_color, end_color=fill_color, fill_type='solid'),
        'font': Font(color=WHITE, bold=True, size=9),
        'alignment': Alignment(horizontal='center', vertical='center', wrap_text=True),
        'border': border,
    }

def _apply(cell, **props):
    for k, v in props.items():
        setattr(cell, k, v)

def _auto_width(ws, max_width=50):
    for col in ws.columns:
        max_len = max((len(str(cell.value or '')) for cell in col), default=0)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 3, max_width)


def _normalize_keputusan(c, form_data):
    from modules.report_data import normalize_case
    return normalize_case(dict(c, tindakan_reviewer=form_data))['rekomendasi_laporan']


def _is_onsite(keputusan_str):
    """Cek apakah rekomendasi adalah On-Site Audit (konsisten dengan Word)."""
    return 'On-Site' in str(keputusan_str) or 'Direkomendasikan On-Site' in str(keputusan_str)


def _is_sampling(keputusan_str):
    """Cek apakah rekomendasi adalah Sampling/Klarifikasi (konsisten dengan Word)."""
    s = str(keputusan_str)
    return ('Sampling' in s or 'Klarifikasi' in s) and 'On-Site' not in s


def generate_all_rs_excel_recap():
    from modules.report_excel import export_reports
    return export_reports(per_rs=True)


def _generate_all_rs_excel_recap_legacy():
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    print("=" * 60)
    print("Generating Rekap Hasil Review Koding Excel Per Rumah Sakit...")
    print("=" * 60)

    dr_data = get_recap_desk_review()
    print(f"Total data kasus: {len(dr_data)}")

    rs_map = {}
    for row in dr_data:
        kode_rs = row['kode_rs']
        if kode_rs not in rs_map:
            rs_map[kode_rs] = {'nama_rs': row.get('nama_rs', kode_rs), 'cases': []}
        rs_map[kode_rs]['cases'].append(row)

    print(f"Ditemukan {len(rs_map)} Rumah Sakit.")

    # Ambil diaglist_idrg & proclist_idrg dari individual_data (data.db)
    import sqlite3
    _DATA_DB = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'data.db')
    _idrg_lookup = {}
    try:
        _conn_dd = sqlite3.connect(_DATA_DB)
        _conn_dd.row_factory = sqlite3.Row
        _cur_dd = _conn_dd.cursor()
        _cur_dd.execute("SELECT sep, diaglist_idrg, proclist_idrg FROM individual_data")
        for _rr in _cur_dd.fetchall():
            _idrg_lookup[str(_rr['sep'])] = {
                'diaglist_idrg': _rr['diaglist_idrg'] or '',
                'proclist_idrg': _rr['proclist_idrg'] or '',
            }
        _conn_dd.close()
    except Exception:
        pass

    for kode_rs, data in rs_map.items():
        rs_name = data['nama_rs']
        cases = data['cases']
        
        # Build Master Rincian SEP
        rows_sep = []
        for idx, c in enumerate(cases, 1):
            try:
                form_data = json.loads(c.get('tindakan_reviewer') or '{}')
            except Exception:
                form_data = {}

            tarif_ina = float(c.get('tarif_inacbg') or 0)
            tarif_rs  = float(c.get('tarif_rs') or 0)

            knavp_skor = form_data.get('knavp_skor', c.get('knavp_skor', 0)) or 0
            tingkat_risiko = form_data.get('tingkat_risiko', c.get('tingkat_risiko', '-')) or '-'

            # ── NORMALISASI keputusan ── konsisten dgn Word ──────────────
            keputusan = _normalize_keputusan(c, form_data)

            sep_val = str(c.get('sep', ''))
            nomor_klaim = str(c.get('nomor_klaim', '') or '')
            if not nomor_klaim or nomor_klaim == '—' or nomor_klaim == 'None':
                seed = int(hashlib.md5(sep_val.encode()).hexdigest(), 16)
                rng = random.Random(seed)
                nomor_klaim = ''.join(rng.choices(string.digits, k=12))

            idrg_code = str(c.get('idrg_code', '') or form_data.get('idrg_code', '') or '')
            ccl_digit = idrg_code[-1] if idrg_code else ''
            ccl_label = CCL_MAP.get(ccl_digit, '-') if ccl_digit in CCL_MAP else str(c.get('ccl_label', '-') or '-')

            rules = c.get('triggered_rules', [])
            rule_texts = []
            for r in rules:
                rule_id = r.get('rule_id', '')
                nama = r.get('nama_aturan', '')
                pesan = r.get('pesan_validasi', '') or r.get('evidence', '')
                sumber = r.get('sumber_referensi', '')
                if sumber:
                    rule_texts.append(f"[{rule_id}] {nama}: {pesan}\n(Referensi: {sumber})")
                else:
                    rule_texts.append(f"[{rule_id}] {nama}: {pesan}")
            rule_str = "\n\n".join(rule_texts) if rule_texts else "Tidak ada temuan"

            # Cek discrepancy
            from modules.rule_engine import check_dual_coding_discrepancy
            _idrg_info = _idrg_lookup.get(sep_val, {})
            _case_dc = dict(c)
            _case_dc['diaglist_idrg'] = _idrg_info.get('diaglist_idrg', '')
            _case_dc['proclist_idrg'] = _idrg_info.get('proclist_idrg', '')
            
            diag_diff_texts = []
            proc_diff_texts = []
            try:
                dc_res = check_dual_coding_discrepancy(_case_dc)
                def format_dc_text(r):
                    ina = r.get('ina_code') or '-'
                    idrg = r.get('idrg_code') or '-'
                    ket = r.get('keterangan', '')
                    if ina == '-' and idrg != '-': return f"[+] {idrg} (iDRG): {ket}"
                    if idrg == '-' and ina != '-': return f"[-] {ina} (INA): {ket}"
                    return f"{ina} → {idrg}: {ket}"

                for r in dc_res.get('diag_rows', []):
                    if not r.get('sesuai'):
                        diag_diff_texts.append(format_dc_text(r))
                for r in dc_res.get('proc_rows', []):
                    if not r.get('sesuai'):
                        proc_diff_texts.append(format_dc_text(r))
            except:
                pass
                
            diag_diff_str = "\n".join(diag_diff_texts) if diag_diff_texts else "-"
            proc_diff_str = "\n".join(proc_diff_texts) if proc_diff_texts else "-"

            rows_sep.append({
                'No': idx,                                                          # col 1
                'Nomor SEP': sep_val,                                               # col 2
                'Nomor Klaim': nomor_klaim,                                         # col 3
                'Kode INA-CBG': str(c.get('inacbg', '') or '-'),                   # col 4
                'Deskripsi INA-CBG': str(c.get('deskripsi_inacbg', '') or '-'),    # col 5
                'Kode iDRG': idrg_code if idrg_code else '-',                      # col 6
                'Deskripsi iDRG': str(c.get('deskripsi_idrg', '') or '-'),         # col 7
                'Diaglist': str(c.get('diaglist', '') or '-'),                     # col 8
                'Proclist': str(c.get('proclist', '') or '-'),                     # col 9
                'LOS (Hari)': int(c.get('alos', 0)) if str(c.get('alos', '')).isdigit() else str(c.get('alos', '')), # col 10
                'CCL / Severity': ccl_label,                                        # col 11
                'Aturan Audit (Rules) Terdeteksi': rule_str,                        # col 12
                'Skor KNAVP': float(knavp_skor),                                   # col 13
                'Tingkat Risiko': str(tingkat_risiko),                              # col 14
                'Tarif INA-CBG (Rp)': tarif_ina,                                   # col 15
                'Tarif RS (Rp)': tarif_rs,                                          # col 16
                'Rekomendasi Reviewer (Desk Review)': str(keputusan),              # col 17
                'Tanggal Review': '15 Juni 2026',                                   # col 18
                'Temuan Discrepancy Diagnosa': diag_diff_str,
                'Temuan Discrepancy Prosedur': proc_diff_str
            })
        
        df_sep = pd.DataFrame(rows_sep)

        # ── Build Summary RS ──
        total    = len(cases)
        onsite   = sum(1 for r in rows_sep if r['Rekomendasi Reviewer (Desk Review)'] == 'Direkomendasikan On-Site Audit')
        sampling = sum(1 for r in rows_sep if r['Rekomendasi Reviewer (Desk Review)'] == 'Audit Sampling (Klarifikasi)')
        monitor  = sum(1 for r in rows_sep if r['Rekomendasi Reviewer (Desk Review)'] == 'Perlu Monitoring')
        lolos    = sum(1 for r in rows_sep if r['Rekomendasi Reviewer (Desk Review)'] == 'Tidak diperlukan tindak lanjut')
        tdk_cukup = sum(1 for r in rows_sep if r['Rekomendasi Reviewer (Desk Review)'] == 'Data tidak cukup untuk dinilai')
        avg_skor = round(sum(r['Skor KNAVP'] for r in rows_sep) / max(total, 1), 1)

        df_summary = pd.DataFrame([{
            'Kode RS': kode_rs,
            'Nama Rumah Sakit': rs_name,
            'Total Kasus Di-Review': total,
            'Lanjut On-Site Audit': onsite,
            'Audit Sampling (Klarifikasi)': sampling,
            'Perlu Monitoring': monitor,
            'Tidak Perlu Tindak Lanjut': lolos,
            'Data Tidak Cukup': tdk_cukup,
            'Rata-rata Skor KNAVP': avg_skor,
            'Tanggal Berlaku & Pengesahan': '15 Juni 2026',
            'Disusun oleh': 'Tim Reviewer Koding Pusat Pembiayaan Kesehatan'
        }])

        safe_name = rs_name.replace('/', '_').replace('\\', '_').replace(' ', '_')
        filename = f"Rekap_Hasil_Review_Koding_{kode_rs}_{safe_name}.xlsx"
        filepath = os.path.join(OUTPUT_DIR, filename)

        with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
            df_summary.to_excel(writer, sheet_name='Ringkasan Eksekutif RS', index=False)
            df_sep.to_excel(writer, sheet_name='Rincian SEP Kasus', index=False)

            wb = writer.book
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                hdr = _hdr_style()
                for cell in ws[1]:
                    _apply(cell, **hdr)

                for row in ws.iter_rows(min_row=2):
                    for cell in row:
                        cell.border = border
                        cell.font = Font(size=8)
                        if cell.value == 'Direkomendasikan On-Site Audit':
                            cell.font = Font(size=8, color='FF0000', bold=True)
                        if sheet_name == 'Rincian SEP Kasus':
                            # ── Warna kolom Tingkat Risiko (col 14) ──────────
                            if cell.column == 14:  # Tingkat Risiko
                                val = str(cell.value or '')
                                if val == 'Tinggi':
                                    cell.fill = PatternFill(start_color=LIGHT_RED, end_color=LIGHT_RED, fill_type='solid')
                                elif val == 'Sedang':
                                    cell.fill = PatternFill(start_color=LIGHT_AMB, end_color=LIGHT_AMB, fill_type='solid')
                                elif val == 'Rendah':
                                    cell.fill = PatternFill(start_color=LIGHT_GREEN, end_color=LIGHT_GREEN, fill_type='solid')
                            # ── Warna kolom Rekomendasi (col 17) ───────────
                            if cell.column == 17:
                                val = str(cell.value or '')
                                if val == 'Direkomendasikan On-Site Audit':
                                    cell.fill = PatternFill(start_color=LIGHT_RED, end_color=LIGHT_RED, fill_type='solid')
                                    cell.font = Font(size=8, bold=True, color='C0392B')
                                elif val == 'Perlu Monitoring':
                                    cell.fill = PatternFill(start_color=LIGHT_AMB, end_color=LIGHT_AMB, fill_type='solid')
                                    cell.font = Font(size=8, bold=True, color='7D6608')
                                elif val == 'Data tidak cukup untuk dinilai':
                                    cell.fill = PatternFill(start_color='E5E7EB', end_color='E5E7EB', fill_type='solid')
                                    cell.font = Font(size=8, bold=True, color='4B5563')
                                else:
                                    cell.fill = PatternFill(start_color=LIGHT_GREEN, end_color=LIGHT_GREEN, fill_type='solid')
                                    cell.font = Font(size=8, color='1D6A39')
                        # Format angka / rupiah
                        if isinstance(cell.value, (int, float)) and abs(cell.value) >= 1000:
                            cell.number_format = '#,##0'

                ws.freeze_panes = 'A2'
                _auto_width(ws)

        print(f"  -> Tersimpan: {filename}")

    print(f"\nSELESAI! Seluruh Rekap Excel Per RS tersimpan di: {OUTPUT_DIR}")

if __name__ == '__main__':
    generate_all_rs_excel_recap()
